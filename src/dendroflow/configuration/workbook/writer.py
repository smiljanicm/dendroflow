"""Write empty and populated CONFIG workbooks using the versioned contract."""

import json
import math
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from uuid import uuid4

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from .preparation import WorkbookExport
from .schema import (
    FORMAT_VERSION,
    GUIDE_SHEET,
    INFO_HEADERS,
    INFO_KEYS,
    INFO_SHEET,
    RESOURCE_SHEETS,
    ROW_ROLES,
    CellKind,
    SheetSpec,
)

_TEXT_KINDS = {
    CellKind.TEXT, CellKind.ID, CellKind.TIMESTAMP,
    CellKind.JSON, CellKind.REFERENCE,
}
_LAST_EXCEL_ROW = 1_048_576
_HEADER_FILL = PatternFill("solid", fgColor="234E52")
_HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF")


def _write_text(sheet: Worksheet, row: int, column: int, value: str) -> None:
    """Write literal text, including strings that Excel could read as formulas."""
    if len(value) > 32767:
        raise ValueError(f"{sheet.title}!{get_column_letter(column)}{row}: text is too long")
    cell = sheet.cell(row, column, value)
    cell.data_type = "s"
    cell.number_format = "@"


def _format_headers(sheet: Worksheet) -> None:
    for cell in sheet[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 32
    sheet.freeze_panes = "A2"


def _add_dropdown(sheet: Worksheet, column: int, choices: tuple[str, ...]) -> None:
    letter = get_column_letter(column)
    validation = DataValidation(
        type="list",
        formula1='"' + ",".join(choices) + '"',
        allow_blank=True,
        showDropDown=False,
        showErrorMessage=True,
        errorStyle="stop",
        errorTitle="Choose a listed value",
        error="Use one of the values in the dropdown.",
    )
    validation.add(f"{letter}2:{letter}{_LAST_EXCEL_ROW}")
    sheet.add_data_validation(validation)


def _add_resource_sheet(workbook: Workbook, spec: SheetSpec) -> Worksheet:
    sheet = workbook.create_sheet(spec.name)
    sheet.append(spec.headers)
    _format_headers(sheet)

    for index, column in enumerate(spec.columns, start=1):
        letter = get_column_letter(index)
        width = max(18, min(34, len(column.name) + 3))
        if column.name in {"description", "path", "reader_options"}:
            width = 48
        sheet.column_dimensions[letter].width = width
        if column.kind in _TEXT_KINDS or column.kind == CellKind.NUMBER:
            # Column defaults help cells created later in Excel. Explicitly
            # format the first empty input row as well, without adding values.
            sheet.column_dimensions[letter].number_format = "@"
            sheet.cell(2, index).number_format = "@"
        if column.name == "row_role":
            _add_dropdown(sheet, index, ROW_ROLES)
        elif column.kind == CellKind.BOOLEAN:
            _add_dropdown(sheet, index, ("TRUE", "FALSE"))
        elif column.name == "reader_type":
            _add_dropdown(sheet, index, ("csv",))
    return sheet


def _add_guide(workbook: Workbook, target_environment: str, *, template: bool) -> None:
    sheet = workbook.active
    sheet.title = GUIDE_SHEET
    sheet.column_dimensions["A"].width = 110
    lines = (
        "DendroFlow control workbook",
        f"Target environment: {target_environment}",
        (
            "This empty template starts a new setup. For existing records, begin "
            "each editing session with a fresh database export."
            if template else
            "This workbook contains exported configuration. Retain existing IDs. "
            "Edit supported fields on edit rows; reference rows supply shared context. "
            "Start the next editing session with a fresh export."
        ),
        (
            "Enter records below row 1. Keep every resource sheet and its headers. "
            "Do not add totals, formulas, or explanatory rows inside resource tables."
        ),
        (
            "New rows: leave the database ID blank, enter a unique ref within the "
            "sheet, and choose row_role=edit. Do not invent database IDs."
        ),
        (
            "Relationship columns use the target sheet's ref. For example, "
            "sensors.sensor_model refers to sensor_models.ref."
        ),
        (
            "Fill required values. A blank nullable value requests null; it does "
            "not mean leave unchanged. Missing rows never request deletion."
        ),
        (
            "Use TRUE or FALSE for boolean fields. Enter identifiers as text "
            "to preserve leading zeros. Paste values without replacing cell formats."
        ),
        (
            "Coordinates, height and azimuth use decimal text to preserve precision. "
            "Keep the text format and use a dot as the decimal separator. "
            "Do not convert exported values to Excel numbers."
        ),
        (
            "Use timestamp text with an explicit offset, for example "
            "2026-01-01T00:00:00Z. Do not use timezone-free Excel date values."
        ),
        (
            "Each new location needs one row in location_labels marked "
            "is_initial=TRUE. Use FALSE for additional labels."
        ),
        (
            "For files, use reader_type=csv and a JSON object in reader_options, "
            "such as {}. Every new file needs at least one interface."
        ),
        (
            "Version 1 cannot update existing files, interfaces, location labels, "
            "or a site's parent. Dropdowns and formatting are editing aids only."
        ),
        (
            "Saving this workbook does not apply changes to databases. "
            "Workbook import and change planning are implemented in later Phase G steps."
        ),
    )
    for row, value in enumerate(lines, start=1):
        _write_text(sheet, row, 1, value)
        sheet.cell(row, 1).alignment = Alignment(wrap_text=True, vertical="center")
        sheet.row_dimensions[row].height = 42
    _format_headers(sheet)


def _destination(path: str | Path, target_environment: str) -> tuple[Path, str]:
    if not isinstance(target_environment, str):
        raise TypeError("target_environment must be a string")
    environment = target_environment.strip()
    if not environment or len(environment) > 128:
        raise ValueError("target_environment must contain 1 to 128 characters")
    if any(ord(character) < 32 for character in environment):
        raise ValueError("target_environment must not contain control characters")

    destination = Path(path)
    if destination.suffix.lower() != ".xlsx":
        raise ValueError("workbook path must end with .xlsx")
    return destination, environment


def _check_export(export: WorkbookExport) -> None:
    """Check the writer's structural boundary, not CONFIG plan applicability."""
    if not isinstance(export, WorkbookExport):
        raise TypeError("export must be a WorkbookExport from prepare_workbook_export")
    if export.scope not in {"all", "sites"}:
        raise ValueError("export scope must be all or sites")
    if (
        not isinstance(export.site_ids, tuple)
        or any(type(key) is not int or not 0 < key < 2**63 for key in export.site_ids)
        or tuple(sorted(set(export.site_ids))) != export.site_ids
        or (export.scope == "all" and export.site_ids)
        or (export.scope == "sites" and not export.site_ids)
    ):
        raise ValueError("site_ids do not match export scope")
    if set(export.frames) != {spec.name for spec in RESOURCE_SHEETS}:
        raise ValueError("export must contain exactly the eleven resource frames")
    for spec in RESOURCE_SHEETS:
        frame = export.frames[spec.name]
        if not frame.columns.is_unique or set(frame.columns) != set(spec.headers):
            raise ValueError(f"{spec.name}: invalid export columns")
        if len(frame) >= _LAST_EXCEL_ROW:
            raise ValueError(f"{spec.name}: too many records for one Excel sheet")
        ids = frame[spec.id_column].tolist()
        if any(
            not isinstance(key, str) or not key.isascii() or not key.isdecimal()
            or len(key) > 19 or not 0 < int(key) < 2**63 or str(int(key)) != key
            for key in ids
        ):
            raise ValueError(f"{spec.name}: IDs must be positive BIGINT decimal text")
        if len(set(ids)) != len(ids) or not frame["ref"].is_unique:
            raise ValueError(f"{spec.name}: duplicate IDs or references")
        if any(role not in ROW_ROLES for role in frame["row_role"]):
            raise ValueError(f"{spec.name}: invalid row_role")
        if export.scope == "all" and any(role != "edit" for role in frame["row_role"]):
            raise ValueError(f"{spec.name}: all-scope rows must have role edit")
    if not set(map(str, export.site_ids)) <= set(export.frames["sites"]["site_id"]):
        raise ValueError("selected site IDs are missing from export")


def _add_records(sheet: Worksheet, spec: SheetSpec, export: WorkbookExport) -> None:
    frame = export.frames[spec.name].loc[:, list(spec.headers)]
    for row_index, values in enumerate(frame.itertuples(index=False, name=None), start=2):
        for column_index, (column, value) in enumerate(zip(spec.columns, values), start=1):
            location = f"{spec.name}!{get_column_letter(column_index)}{row_index}"
            if value is None:
                if not column.nullable:
                    raise ValueError(f"{location}: required value is null")
            elif column.kind == CellKind.BOOLEAN:
                if type(value) is not bool:
                    raise ValueError(f"{location}: expected a boolean")
                sheet.cell(row_index, column_index, value)
            elif column.kind == CellKind.NUMBER:
                if type(value) not in {int, float} or not math.isfinite(value):
                    raise ValueError(f"{location}: expected a finite number")
                # repr(float) round-trips to the same Python float. Excel's
                # numeric precision would not preserve every stored DB value.
                _write_text(sheet, row_index, column_index, repr(value))
            else:
                if not isinstance(value, str) or not value or value != value.strip():
                    raise ValueError(f"{location}: expected nonblank prepared text")
                _write_text(sheet, row_index, column_index, value)
        if frame.iloc[row_index - 2]["row_role"] == "reference":
            for cell in sheet[row_index]:
                cell.fill = PatternFill("solid", fgColor="E8EDF0")
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(spec.headers))}{max(1, len(frame) + 1)}"


def _write_workbook(
    destination: Path, environment: str,
    *, export: WorkbookExport | None, exported_at: str | None,
) -> Path:

    workbook = Workbook()
    try:
        _add_guide(workbook, environment, template=export is None)
        info = workbook.create_sheet(INFO_SHEET)
        info.append(INFO_HEADERS)
        values = {
            "format_version": FORMAT_VERSION,
            "workbook_id": str(uuid4()),
            "target_environment": environment,
            "exported_at": exported_at,
            "scope": "template" if export is None else export.scope,
            "site_ids": json.dumps([] if export is None else list(map(str, export.site_ids))),
        }
        for row, key in enumerate(INFO_KEYS, start=2):
            _write_text(info, row, 1, key)
            value = values[key]
            if isinstance(value, str):
                _write_text(info, row, 2, value)
            else:
                info.cell(row, 2, value)
        _format_headers(info)
        info.column_dimensions["A"].width = 24
        info.column_dimensions["B"].width = 52

        for spec in RESOURCE_SHEETS:
            sheet = _add_resource_sheet(workbook, spec)
            if export is not None:
                _add_records(sheet, spec, export)

        with BytesIO() as buffer:
            workbook.save(buffer)
            with destination.open("xb") as output:
                output.write(buffer.getvalue())
    finally:
        workbook.close()

    return destination


def write_workbook_template(
    path: str | Path, *, target_environment: str,
) -> Path:
    """Save a new empty XLSX template without database access or overwriting.

    Parent directories must exist. Serialization finishes in memory before the
    output path is opened. Filesystem write failures can leave a partial file.
    """
    destination, environment = _destination(path, target_environment)
    return _write_workbook(destination, environment, export=None, exported_at=None)


def write_workbook_export(
    path: str | Path, export: WorkbookExport, *, target_environment: str,
    exported_at: datetime | None = None,
) -> Path:
    """Save prepared frames as a new populated XLSX workbook.

    Pass the result of prepare_workbook_export; this writer does not recheck
    database state or full configuration semantics. Structural checks account
    for the mutable frames. There is no database access or automatic overwrite.

    exported_at defaults to file-generation time, not a database snapshot time.
    A supplied timestamp must be aware and is written in UTC. It does not prove
    freshness or consistency across databases. Environment is a caller-supplied
    non-secret label, not verified connection identity. As with templates, all
    serialization occurs before opening the output, but saving is not atomic
    in the event of a filesystem write failure.
    """
    destination, environment = _destination(path, target_environment)
    _check_export(export)
    when = datetime.now(timezone.utc) if exported_at is None else exported_at
    if not isinstance(when, datetime) or when.utcoffset() is None:
        raise ValueError("exported_at must be a timezone-aware datetime")
    timestamp = when.astimezone(timezone.utc).isoformat().replace("+00:00", "Z")
    return _write_workbook(destination, environment, export=export, exported_at=timestamp)
