"""Read workbook structure and metadata without interpreting resource values."""

import json
import re
from dataclasses import dataclass
from datetime import datetime, timedelta
from pathlib import Path
from typing import NoReturn
from uuid import UUID
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.worksheet import Worksheet

from .schema import (
    EXPORT_SCOPES,
    FORMAT_VERSION,
    GUIDE_SHEET,
    INFO_KEYS,
    INFO_SHEET,
    RESOURCE_SHEETS,
    collect_header_issues,
)


@dataclass(frozen=True)
class WorkbookIssue:
    sheet: str | None
    coordinate: str | None
    message: str


class WorkbookReadError(ValueError):
    """A workbook could not be read; no partial document is returned."""

    def __init__(self, issues: tuple[WorkbookIssue, ...]):
        self.issues = issues
        super().__init__("; ".join(
            f"{issue.sheet or 'workbook'}"
            f"{'!' + issue.coordinate if issue.coordinate else ''}: {issue.message}"
            for issue in issues
        ))


@dataclass(frozen=True)
class WorkbookCell:
    sheet: str
    coordinate: str
    value: object
    data_type: str


@dataclass(frozen=True)
class WorkbookMetadata:
    format_version: int
    workbook_id: UUID
    target_environment: str
    exported_at: datetime | None
    scope: str
    site_ids: tuple[int, ...]


@dataclass(frozen=True)
class WorkbookDocument:
    """Raw resource rows keyed by header name, retaining physical cell locations.

    Sheets and cells use schema order; rows keep workbook order. The dictionaries
    are caller-owned. No resource value or relationship validation is implied.
    """

    metadata: WorkbookMetadata
    sheets: dict[str, tuple[dict[str, WorkbookCell], ...]]


def _blank(value: object) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _fail(cell: WorkbookCell, message: str) -> NoReturn:
    raise WorkbookReadError((WorkbookIssue(cell.sheet, cell.coordinate, message),))


def _rows(workbook) -> dict[str, tuple[dict[str, WorkbookCell], ...]]:
    """Check exact headers before mapping cells so duplicates cannot disappear."""
    headers = {}
    issues = []
    for name in workbook.sheetnames:
        if name == GUIDE_SHEET:
            continue
        sheet = workbook[name]
        if not isinstance(sheet, Worksheet):
            issues.append(WorkbookIssue(name, None, "expected a worksheet, not a chart sheet"))
            continue
        for merged in sheet.merged_cells.ranges:
            issues.append(WorkbookIssue(sheet.title, str(merged), "merged cells are not supported"))
        cells = list(sheet[1])
        while cells and _blank(cells[-1].value):
            cells.pop()
        names = []
        for cell in cells:
            if cell.data_type in {"f", "e"} or not isinstance(cell.value, str) or _blank(cell.value):
                issues.append(WorkbookIssue(
                    sheet.title, cell.coordinate, "header must be nonblank literal text",
                ))
            else:
                names.append(cell.value)
        headers[sheet.title] = names

    for issue in collect_header_issues(headers):
        coordinate = None
        if issue.column and issue.sheet in headers:
            matches = [cell.coordinate for cell in workbook[issue.sheet][1]
                       if cell.value == issue.column]
            coordinate = matches[-1] if matches else None
        issues.append(WorkbookIssue(issue.sheet, coordinate, (
            f"{issue.message}: {issue.column}" if issue.column else issue.message
        )))
    if issues:
        raise WorkbookReadError(tuple(issues))

    tables = {}
    for name, columns in headers.items():
        records = []
        for cells in workbook[name].iter_rows(min_row=2):
            if all(_blank(cell.value) for cell in cells):
                continue
            for cell in cells[len(columns):]:
                if not _blank(cell.value):
                    issues.append(WorkbookIssue(name, cell.coordinate, "value has no column header"))
            records.append({
                column: WorkbookCell(name, cell.coordinate, cell.value, cell.data_type)
                for column, cell in zip(columns, cells)
            })
        tables[name] = tuple(records)
    if issues:
        raise WorkbookReadError(tuple(issues))
    return tables


def _metadata(rows, expected_environment: str) -> WorkbookMetadata:
    cells = {}
    issues = []
    for row in rows:
        key, value = row["key"], row["value"]
        if key.data_type in {"f", "e"} or not isinstance(key.value, str) or key.value not in INFO_KEYS:
            issues.append(WorkbookIssue(INFO_SHEET, key.coordinate, "unknown or invalid metadata key"))
        elif key.value in cells:
            issues.append(WorkbookIssue(INFO_SHEET, key.coordinate, f"duplicate key: {key.value}"))
        else:
            cells[key.value] = value
        if value.data_type in {"f", "e"}:
            issues.append(WorkbookIssue(INFO_SHEET, value.coordinate, "metadata cannot be a formula or Excel error"))
    for key in INFO_KEYS:
        if key not in cells:
            issues.append(WorkbookIssue(INFO_SHEET, None, f"missing key: {key}"))
    if issues:
        raise WorkbookReadError(tuple(issues))

    def text(key):
        cell = cells[key]
        if not isinstance(cell.value, str) or _blank(cell.value) or cell.value != cell.value.strip():
            _fail(cell, f"{key} must be nonblank text without surrounding whitespace")
        return cell.value

    version = cells["format_version"]
    if type(version.value) is not int or version.value != FORMAT_VERSION:
        _fail(version, f"format_version must be integer {FORMAT_VERSION}")
    identifier = text("workbook_id")
    try:
        workbook_id = UUID(identifier)
    except ValueError:
        _fail(cells["workbook_id"], "workbook_id must be a UUID")

    environment = text("target_environment")
    if len(environment) > 128 or any(ord(character) < 32 for character in environment):
        _fail(cells["target_environment"], "invalid target_environment label")
    if environment != expected_environment:
        _fail(cells["target_environment"], "target_environment does not match the expected environment")

    scope = text("scope")
    if scope not in EXPORT_SCOPES:
        _fail(cells["scope"], "scope must be all, sites or template")
    try:
        selected = json.loads(text("site_ids"))
    except json.JSONDecodeError:
        _fail(cells["site_ids"], "site_ids must be a JSON array of decimal strings")
    if not isinstance(selected, list) or any(
        not isinstance(key, str) or not key.isascii() or not key.isdecimal()
        or len(key) > 19 or not 0 < int(key) < 2**63 or str(int(key)) != key
        for key in selected
    ):
        _fail(cells["site_ids"], "site_ids must contain positive BIGINT decimal strings")
    if len(set(selected)) != len(selected):
        _fail(cells["site_ids"], "site_ids must not contain duplicates")
    if (scope == "sites") != bool(selected):
        _fail(cells["site_ids"], "site_ids must be nonempty only for sites scope")

    exported_at = None
    timestamp = cells["exported_at"]
    if scope == "template":
        if not _blank(timestamp.value):
            _fail(timestamp, "template exported_at must be blank")
    else:
        value = text("exported_at")
        if not re.fullmatch(r"\d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2}(?:\.\d{1,6})?(?:Z|\+00:00)", value):
            _fail(timestamp, "exported_at must be ISO 8601 UTC text")
        try:
            exported_at = datetime.fromisoformat(value.replace("Z", "+00:00"))
        except ValueError:
            _fail(timestamp, "exported_at must be a valid UTC datetime")
        if exported_at.utcoffset() != timedelta(0):
            _fail(timestamp, "exported_at must be UTC")
    return WorkbookMetadata(
        FORMAT_VERSION, workbook_id, environment, exported_at, scope,
        tuple(sorted(map(int, selected))),
    )


def read_workbook(path: str | Path, *, expected_environment: str) -> WorkbookDocument:
    """Load an XLSX and check structure/provenance without database access.

    Resource cells are raw, including formulas/errors for G3.b to reject. A
    successful read does not authorize comparison or apply. Guide content and
    wholly blank rows are ignored; hidden records are read. Sheet/column order
    can vary. The expected environment is a caller-supplied label, not a verified
    connection identity. Input files are never modified.
    """
    if (
        not isinstance(expected_environment, str) or not expected_environment
        or expected_environment != expected_environment.strip()
        or len(expected_environment) > 128
        or any(ord(character) < 32 for character in expected_environment)
    ):
        raise ValueError("expected_environment must be a nonblank label of 1 to 128 characters")
    source = Path(path).expanduser()
    if source.suffix.lower() != ".xlsx":
        raise WorkbookReadError((WorkbookIssue(None, None, "input path must end with .xlsx"),))
    try:
        stream = source.open("rb")
    except OSError as error:
        raise WorkbookReadError((WorkbookIssue(None, None, f"cannot open input: {error}"),)) from error
    with stream:
        try:
            workbook = load_workbook(stream, data_only=False, keep_links=False)
        except (BadZipFile, InvalidFileException, OSError, ValueError, KeyError, SyntaxError) as error:
            raise WorkbookReadError((WorkbookIssue(None, None, "cannot read XLSX workbook"),)) from error
        try:
            tables = _rows(workbook)
            metadata = _metadata(tables[INFO_SHEET], expected_environment)
            return WorkbookDocument(metadata, {
                spec.name: tuple({column: row[column] for column in spec.headers} for row in tables[spec.name])
                for spec in RESOURCE_SHEETS
            })
        finally:
            workbook.close()
