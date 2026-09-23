from datetime import datetime, timezone
from uuid import UUID
from zipfile import ZipFile

import pandas as pd
import pytest
from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font

from dendroflow.configuration.workbook.preparation import prepare_workbook_export
from dendroflow.configuration.workbook.reader import WorkbookReadError, read_workbook
from dendroflow.configuration.workbook.schema import INFO_KEYS, RESOURCE_SHEETS
from dendroflow.configuration.workbook.source import SOURCE_COLUMNS
from dendroflow.configuration.workbook.writer import (
    write_workbook_export,
    write_workbook_template,
)


@pytest.fixture
def template(tmp_path):
    return write_workbook_template(tmp_path / "control.xlsx", target_environment="local-dev")


def edit(path, change):
    workbook = load_workbook(path)
    try:
        change(workbook)
        workbook.save(path)
    finally:
        workbook.close()


def set_info(workbook, key, value):
    for row in workbook["workbook_info"].iter_rows(min_row=2):
        if row[0].value == key:
            row[1].value = value
            return
    raise AssertionError(key)


def read(path):
    return read_workbook(path, expected_environment="local-dev")


def test_empty_template_has_metadata_and_all_empty_tables(template):
    document = read(template)
    assert document.metadata.format_version == 1
    assert isinstance(document.metadata.workbook_id, UUID)
    assert document.metadata.target_environment == "local-dev"
    assert document.metadata.scope == "template"
    assert document.metadata.exported_at is None
    assert document.metadata.site_ids == ()
    assert tuple(document.sheets) == tuple(spec.name for spec in RESOURCE_SHEETS)
    assert all(rows == () for rows in document.sheets.values())


@pytest.mark.parametrize("site_ids", [None, [2]])
def test_exported_workbook_reads_metadata_and_raw_values(tmp_path, site_ids):
    records = {"sites": [
        (1, "region", "Region", None, None, None, None),
        (2, "000001", "NA", None, 53.123456789012344, 0.0, 1),
    ], "variables": [(3, "level", False, None)]}
    frames = {
        name: pd.DataFrame(records.get(name, []), columns=columns, dtype=object)
        for name, columns in SOURCE_COLUMNS.items()
    }
    export = prepare_workbook_export(frames, site_ids=site_ids)
    when = datetime(2026, 1, 1, 0, 0, 0, 123456, tzinfo=timezone.utc)
    path = write_workbook_export(
        tmp_path / "export.xlsx", export, target_environment="local-dev", exported_at=when,
    )
    document = read(path)
    assert document.metadata.scope == export.scope
    assert document.metadata.site_ids == export.site_ids
    assert document.metadata.exported_at == when
    row = document.sheets["sites"][1]
    assert row["site_id"].value == "2"
    assert row["site_code"].value == "000001"
    assert row["name"].value == "NA"
    assert row["description"].value is None
    assert row["latitude"].value == "53.123456789012344"
    assert row["latitude"].data_type == "s"
    assert row["parent"].value == "site_1"
    if site_ids is None:
        assert document.sheets["variables"][0]["derived"].value is False
        assert document.sheets["variables"][0]["derived"].data_type == "b"
    else:
        assert document.sheets["sites"][0]["row_role"].value == "reference"


def test_reordering_preserves_header_mapping_and_physical_coordinates(template):
    def change(workbook):
        sheet = workbook["sites"]
        sheet.append(["2", "second", "edit", "0002", "Second"])
        sheet.append(["1", "first", "edit", "0001", "First"])
        # Move name to column A; do not sort the resource rows.
        rows = list(sheet.values)
        order = [4, 0, 1, 2, 3, 5, 6, 7, 8]
        for row_number, values in enumerate(rows, start=1):
            for column, index in enumerate(order, start=1):
                sheet.cell(row_number, column).value = values[index]
        info = workbook["workbook_info"]
        info_rows = list(info.values)
        for number, values in enumerate([info_rows[0], *reversed(info_rows[1:])], start=1):
            info.cell(number, 1).value = values[1]
            info.cell(number, 2).value = values[0]
        for name in reversed(workbook.sheetnames):
            workbook.move_sheet(name, offset=-len(workbook.sheetnames))

    edit(template, change)
    document = read(template)
    first, second = document.sheets["sites"]
    assert tuple(first) == RESOURCE_SHEETS[0].headers
    assert first["name"].value == "Second"
    assert first["name"].coordinate == "A3"  # template already had an empty input row
    assert first["name"].sheet == "sites"
    assert second["site_id"].value == "1"
    assert second["site_id"].coordinate == "B4"


def test_blank_rows_and_formatting_are_ignored_but_hidden_rows_are_read(template):
    def change(workbook):
        sheet = workbook["variables"]
        sheet["D4"] = "NA"
        sheet["E4"] = False
        sheet["F4"] = " NULL "
        sheet["A6"] = 0
        sheet["B7"] = "   "
        sheet["B20"].font = Font(bold=True)
        sheet["Z20"].number_format = "@"
        sheet.row_dimensions[4].hidden = True
        sheet.sheet_state = "hidden"

    edit(template, change)
    rows = read(template).sheets["variables"]
    assert len(rows) == 2
    assert rows[0]["variable"].coordinate == "D4"
    assert rows[0]["derived"].value is False
    assert rows[0]["description"].value == " NULL "
    assert rows[1]["variable_id"].value == 0  # G3.b will reject the invalid ID.


@pytest.mark.parametrize("guide", ["removed", "changed"])
def test_guide_is_optional_and_ignored(template, guide):
    def change(workbook):
        if guide == "removed":
            del workbook["guide"]
        else:
            workbook["guide"]["A1"] = "=1+1"
            workbook["guide"].merge_cells("A2:C4")

    edit(template, change)
    assert read(template).metadata.scope == "template"


@pytest.mark.parametrize("defect", ["missing_sheet", "unknown_sheet", "missing_header",
                                     "unknown_header", "duplicate_header", "numeric_header",
                                     "blank_header", "formula_header", "merged_cells", "extra_value"])
def test_invalid_structure_reports_a_location(template, defect):
    def change(workbook):
        sheet = workbook["sites"]
        if defect == "missing_sheet":
            del workbook["sites"]
        elif defect == "unknown_sheet":
            workbook.create_sheet("notes")
        elif defect == "missing_header":
            sheet.delete_cols(5)
        elif defect == "unknown_header":
            sheet["E1"] = "Name"
        elif defect == "duplicate_header":
            sheet["J1"] = "name"
        elif defect == "numeric_header":
            sheet["E1"] = 42
        elif defect == "blank_header":
            sheet["E1"] = None
        elif defect == "formula_header":
            sheet["E1"] = '= "name"'
        elif defect == "merged_cells":
            sheet.merge_cells("F3:G3")
        else:
            sheet["J5"] = "unheaded value"

    edit(template, change)
    with pytest.raises(WorkbookReadError) as caught:
        read(template)
    assert caught.value.issues
    assert all(issue.sheet for issue in caught.value.issues)
    if defect == "extra_value":
        assert caught.value.issues[0].coordinate == "J5"
    if defect == "duplicate_header":
        assert any(issue.coordinate == "J1" for issue in caught.value.issues)


@pytest.mark.parametrize("defect", ["missing", "duplicate", "unknown", "formula_key", "formula_value", "error_value"])
def test_invalid_metadata_structure_is_rejected(template, defect):
    def change(workbook):
        sheet = workbook["workbook_info"]
        if defect == "missing":
            sheet.delete_rows(2)
        elif defect == "duplicate":
            sheet.append(["scope", "template"])
        elif defect == "unknown":
            sheet.append(["note", "extra"])
        elif defect == "formula_key":
            sheet["A2"] = '="format_version"'
        elif defect == "formula_value":
            sheet["B2"] = "=1"
        else:
            sheet["B2"] = "#REF!"

    edit(template, change)
    with pytest.raises(WorkbookReadError) as caught:
        read(template)
    assert all(issue.sheet == "workbook_info" for issue in caught.value.issues)


@pytest.mark.parametrize(("key", "value"), [
    ("format_version", 2), ("format_version", "1"), ("format_version", True),
    ("workbook_id", "invalid"), ("workbook_id", None),
    ("target_environment", "production"), ("target_environment", " local-dev "),
    ("scope", "other"), ("site_ids", "not json"), ("site_ids", "{}"),
    ("site_ids", "[2]"), ("site_ids", '["0"]'), ("site_ids", '["01"]'),
    ("site_ids", '["9223372036854775808"]'),
    ("site_ids", '["2", "2"]'), ("site_ids", '["2"]'),
    ("exported_at", "2026-01-01T00:00:00Z"),
])
def test_invalid_metadata_value_has_cell_diagnostic(template, key, value):
    edit(template, lambda workbook: set_info(workbook, key, value))
    with pytest.raises(WorkbookReadError) as caught:
        read(template)
    issue = caught.value.issues[0]
    assert issue.sheet == "workbook_info"
    assert issue.coordinate == f"B{INFO_KEYS.index(key) + 2}"


@pytest.mark.parametrize("timestamp", [None, "2026-01-01", "2026-01-01T00:00:00",
                                       "2026-01-01T01:00:00+01:00", "2026-02-30T00:00:00Z",
                                       datetime(2026, 1, 1, tzinfo=timezone.utc).replace(tzinfo=None)])
def test_export_requires_valid_utc_timestamp_text(template, timestamp):
    def change(workbook):
        set_info(workbook, "scope", "all")
        set_info(workbook, "exported_at", timestamp)

    edit(template, change)
    with pytest.raises(WorkbookReadError, match="exported_at"):
        read(template)


def test_selected_site_ids_are_canonicalized_without_reading_database(template):
    def change(workbook):
        set_info(workbook, "scope", "sites")
        set_info(workbook, "site_ids", '["9223372036854775807", "2"]')
        set_info(workbook, "exported_at", "2026-01-01T00:00:00+00:00")

    edit(template, change)
    assert read(template).metadata.site_ids == (2, 2**63 - 1)
    # Whether these sites exist or are represented by rows is checked later.


def test_sites_scope_requires_nonempty_selection(template):
    edit(template, lambda workbook: set_info(workbook, "scope", "sites"))
    with pytest.raises(WorkbookReadError, match="site_ids"):
        read(template)


def test_resource_cell_kinds_are_retained_for_later_validation(template):
    def change(workbook):
        sheet = workbook["sites"]
        sheet["A2"] = "0001"
        sheet["D2"] = "=1+1"
        sheet["E2"] = "#REF!"
        sheet["F2"] = "=literal"
        sheet["F2"].data_type = "s"
        sheet["G2"] = 0
        sheet["H2"] = datetime(2026, 1, 1, tzinfo=timezone.utc).replace(tzinfo=None)

    edit(template, change)
    row = read(template).sheets["sites"][0]
    assert row["site_id"].value == "0001"
    assert row["site_code"].value == "=1+1"
    assert row["site_code"].data_type == "f"
    assert row["name"].data_type == "e"
    assert row["description"].data_type == "s"
    assert row["description"].value == "=literal"
    assert row["latitude"].value == 0
    assert row["longitude"].data_type == "d"


@pytest.mark.parametrize("environment", [None, "", " padded ", "x" * 129, "local\ndev"])
def test_expected_environment_is_required_and_validated(template, environment):
    with pytest.raises(ValueError, match="expected_environment"):
        read_workbook(template, expected_environment=environment)


@pytest.mark.parametrize("defect", ["missing", "wrong_extension", "not_zip", "incomplete_zip"])
def test_unreadable_input_raises_workbook_error(tmp_path, defect):
    path = tmp_path / ("bad.xls" if defect == "wrong_extension" else "bad.xlsx")
    if defect in {"wrong_extension", "not_zip"}:
        path.write_bytes(b"not a workbook")
    elif defect == "incomplete_zip":
        with ZipFile(path, "w") as archive:
            archive.writestr("unrelated.txt", "data")
    with pytest.raises(WorkbookReadError):
        read(path)


def test_reading_preserves_input_and_never_connects_to_database(template, monkeypatch):
    def fail_connect(*args, **kwargs):
        pytest.fail("reader attempted a database connection")

    monkeypatch.setattr("dendroflow.database.connect", fail_connect)
    monkeypatch.setattr("dendroflow.configuration.workbook.source.connect", fail_connect)
    original = template.read_bytes()
    assert read(str(template)).metadata.scope == "template"
    assert template.read_bytes() == original
    template.unlink()  # No input handle needs to survive the read.


def test_resource_validation_is_not_claimed_by_structural_read(template):
    def change(workbook):
        sheet = workbook["sensors"]
        sheet["A2"] = "not an ID"
        sheet["C2"] = "not a role"
        sheet["E2"] = "missing_model"

    edit(template, change)
    row = read(template).sheets["sensors"][0]
    assert row["sensor_id"].value == "not an ID"
    assert row["sensor_model"].value == "missing_model"


@pytest.mark.parametrize("name", ["notes", "sites"])
def test_chart_sheets_cannot_replace_or_extend_resource_sheets(template, name):
    def change(workbook):
        if name in workbook.sheetnames:
            del workbook[name]
        chart = BarChart()
        chart.add_data(Reference(workbook["workbook_info"], min_col=2, min_row=2, max_row=2))
        workbook.create_chartsheet(name).add_chart(chart)

    edit(template, change)
    with pytest.raises(WorkbookReadError, match="chart sheet"):
        read(template)


def test_malformed_xml_has_a_workbook_diagnostic(template):
    with ZipFile(template) as archive:
        contents = {name: archive.read(name) for name in archive.namelist()}
    contents["xl/workbook.xml"] = b"<broken>"
    with ZipFile(template, "w") as archive:
        for name, content in contents.items():
            archive.writestr(name, content)
    with pytest.raises(WorkbookReadError, match="cannot read XLSX"):
        read(template)
