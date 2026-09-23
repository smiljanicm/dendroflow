import math
from copy import deepcopy
from dataclasses import replace
from datetime import datetime, timedelta, timezone
from uuid import UUID

import pandas as pd
import pytest
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from dendroflow.configuration.workbook.parsing import WorkbookParseError, parse_workbook
from dendroflow.configuration.workbook.preparation import prepare_workbook_export
from dendroflow.configuration.workbook.reader import (
    WorkbookCell,
    WorkbookDocument,
    WorkbookMetadata,
    read_workbook,
)
from dendroflow.configuration.workbook.schema import (
    RESOURCE_SHEETS,
    CellKind,
    get_sheet_spec,
)
from dendroflow.configuration.workbook.source import SOURCE_COLUMNS
from dendroflow.configuration.workbook.writer import (
    write_workbook_export,
    write_workbook_template,
)

START = datetime(2025, 1, 1, tzinfo=timezone.utc)


@pytest.fixture
def document():
    return WorkbookDocument(
        WorkbookMetadata(1, UUID(int=1), "local-dev", None, "template", ()),
        {spec.name: () for spec in RESOURCE_SHEETS},
    )


def with_cell(document, sheet, column, value, data_type=None):
    spec = get_sheet_spec(sheet)
    defaults = {
        CellKind.TEXT: "example", CellKind.ID: "1", CellKind.NUMBER: "0",
        CellKind.REFERENCE: "target", CellKind.BOOLEAN: False,
        CellKind.TIMESTAMP: "2025-01-01T00:00:00Z", CellKind.JSON: "{}",
    }
    row = {}
    for index, field in enumerate(spec.columns, start=1):
        raw = None if field.nullable else defaults[field.kind]
        if field.name == "row_role":
            raw = "edit"
        if field.name == "reader_type":
            raw = "csv"
        if field.name == column:
            raw = value
        kind = "b" if type(raw) is bool else "n" if type(raw) in {int, float} or raw is None else "s"
        row[field.name] = WorkbookCell(sheet, f"{get_column_letter(index)}7", raw, kind)
    if data_type is not None:
        row[column] = replace(row[column], data_type=data_type)
    return replace(document, sheets={**document.sheets, sheet: (row,)})


def value_from(document, sheet, column, value, data_type=None):
    parsed = parse_workbook(with_cell(document, sheet, column, value, data_type))
    return parsed.frames[sheet].at[0, column]


def test_empty_template_returns_all_headers_and_metadata(tmp_path):
    path = write_workbook_template(tmp_path / "template.xlsx", target_environment="local-dev")
    document = read_workbook(path, expected_environment="local-dev")
    result = parse_workbook(document)
    assert result.metadata == document.metadata
    for spec in RESOURCE_SHEETS:
        assert result.frames[spec.name].empty
        assert tuple(result.frames[spec.name].columns) == spec.headers
        assert tuple(result.coordinates[spec.name].columns) == spec.headers
        assert all(dtype == object for dtype in result.frames[spec.name].dtypes)


def test_populated_export_preserves_values_through_read_and_parse(tmp_path):
    key = 2**63 - 1
    records = {
        "sites": [(key, "000001", "Český les", None, 53.123456789012344, 0.0, None)],
        "location_types": [(10, "well", None)],
        "sensor_types": [(20, "logger", None)],
        "variables": [(30, "level", False, None)],
        "sensor_models": [(40, "Maker", "Model", 20)],
        "sensors": [(50, "000123", 40, "NA")],
        "locations": [(60, key, 10, None, None, -0.0, 0.0)],
        "location_labels": [(70, 60, "Well", START, None)],
        "deployments": [(80, 50, 60, 30, START, None)],
        "files": [(90, "/data/測定.csv", "UTC", "%Y-%m-%d", {
            "reader": "csv", "options": {"skiprows": [0, 2], "na_values": ["NA", ""]},
        })],
        "interfaces": [(100, 90, 80, "time", "level", "cm")],
    }
    frames = {
        name: pd.DataFrame(records.get(name, []), columns=columns, dtype=object)
        for name, columns in SOURCE_COLUMNS.items()
    }
    export = prepare_workbook_export(frames)
    path = write_workbook_export(tmp_path / "export.xlsx", export, target_environment="local-dev")
    parsed = parse_workbook(read_workbook(path, expected_environment="local-dev"))
    assert all(len(frame) == 1 for frame in parsed.frames.values())
    site = parsed.frames["sites"].iloc[0]
    assert type(site.site_id) is int and site.site_id == key
    assert site.site_code == "000001"
    assert site.latitude == 53.123456789012344
    assert site.description is None
    assert parsed.frames["sensors"].at[0, "serial_number"] == "000123"
    assert parsed.frames["variables"].at[0, "derived"] is False
    assert parsed.frames["location_labels"].at[0, "is_initial"] is True
    assert parsed.frames["deployments"].at[0, "valid_from"] == START
    assert type(parsed.frames["deployments"].at[0, "valid_from"]) is datetime
    assert parsed.frames["deployments"].at[0, "valid_to"] is None
    assert math.copysign(1, parsed.frames["locations"].at[0, "height_above_ground"]) == -1
    assert parsed.frames["interfaces"].at[0, "deployment"] == "deployment_80"
    assert parsed.frames["files"].at[0, "reader_options"] == {
        "skiprows": [0, 2], "na_values": ["NA", ""],
    }
    for name, frame in parsed.frames.items():
        assert all(dtype == object for dtype in frame.dtypes)
        assert parsed.coordinates[name].shape == frame.shape


def test_reordered_columns_and_blank_rows_keep_excel_coordinates(tmp_path):
    path = write_workbook_template(tmp_path / "template.xlsx", target_environment="local-dev")
    workbook = load_workbook(path)
    try:
        sheet = workbook["variables"]
        sheet["D5"] = "level"
        sheet["B5"] = "v"
        sheet["C5"] = "edit"
        sheet["E5"] = "FALSE"
        for row in sheet:
            row[0].value, row[4].value = row[4].value, row[0].value
        workbook.save(path)
    finally:
        workbook.close()
    parsed = parse_workbook(read_workbook(path, expected_environment="local-dev"))
    assert parsed.frames["variables"].at[0, "variable_id"] is None
    assert parsed.frames["variables"].at[0, "derived"] is False
    assert parsed.coordinates["variables"].at[0, "derived"] == "A5"
    assert parsed.coordinates["variables"].at[0, "variable_id"] == "E5"


@pytest.mark.parametrize("raw", [None, "", " \t\n "])
def test_nullable_blanks_become_none(document, raw):
    assert value_from(document, "sites", "description", raw) is None


@pytest.mark.parametrize("raw", ["000001", "NA", "NULL", "=literal", "#REF!", "  Name  "])
def test_text_is_trimmed_without_missing_value_inference(document, raw):
    assert value_from(document, "sites", "name", raw, "s") == raw.strip()


@pytest.mark.parametrize("raw", [0, False, 1.2, None, "   ", "bad\x00text"])
def test_invalid_required_text_is_not_coerced(document, raw):
    with pytest.raises(WorkbookParseError, match="sites!E7: name"):
        value_from(document, "sites", "name", raw)


@pytest.mark.parametrize("raw", ["1", "9223372036854775807", " 42 "])
def test_ids_are_exact_python_integers(document, raw):
    result = value_from(document, "sites", "site_id", raw)
    assert type(result) is int and result == int(raw)


@pytest.mark.parametrize("raw", [1, 1.0, True, "0", "-1", "01", "1.0", "1e3", "9223372036854775808"])
def test_invalid_or_numeric_excel_ids_are_rejected(document, raw):
    with pytest.raises(WorkbookParseError, match="BIGINT ID as decimal text"):
        value_from(document, "sites", "site_id", raw)


@pytest.mark.parametrize(("raw", "expected"), [
    (0, 0.0), (1.25, 1.25), (" 1.25 ", 1.25), ("+.5", 0.5), ("1e-3", 0.001),
    ("53.123456789012344", 53.123456789012344), ("5e-324", 5e-324),
])
def test_finite_numbers_accept_decimal_text_and_native_numbers(document, raw, expected):
    result = value_from(document, "sites", "latitude", raw)
    assert type(result) is float and result == expected


@pytest.mark.parametrize("raw", [True, "1,25", "1_000", "NaN", float("nan"), float("inf"), "1e309", "1e-999"])
def test_invalid_nonfinite_and_underflowing_numbers_are_rejected(document, raw):
    with pytest.raises(WorkbookParseError, match="latitude"):
        value_from(document, "sites", "latitude", raw)


@pytest.mark.parametrize(("raw", "expected"), [(True, True), (False, False), (" TRUE ", True), ("false", False)])
def test_boolean_values_are_explicit(document, raw, expected):
    assert value_from(document, "variables", "derived", raw) is expected


@pytest.mark.parametrize("raw", [0, 1, "yes", "0", None])
def test_implicit_boolean_values_are_rejected(document, raw):
    with pytest.raises(WorkbookParseError, match="derived"):
        value_from(document, "variables", "derived", raw)


@pytest.mark.parametrize("raw", ["2025-01-01T00:00:00.123456Z", "2025-01-01T02:00:00.123456+02:00"])
def test_timestamp_offsets_normalize_to_utc_without_losing_microseconds(document, raw):
    result = value_from(document, "deployments", "valid_from", raw)
    assert result == START + timedelta(microseconds=123456)
    assert result.tzinfo is timezone.utc


@pytest.mark.parametrize("raw", [
    "2025-01-01", "2025-01-01T00:00:00", "2025-02-30T00:00:00Z",
    "2025-01-01T00:00:00.1234567Z", "2025-01-01T00:00:00+01:60",
    "0001-01-01T00:00:00+01:00", START, 45658,
])
def test_invalid_or_ambiguous_timestamps_are_rejected(document, raw):
    with pytest.raises(WorkbookParseError, match="valid_from"):
        value_from(document, "deployments", "valid_from", raw)


def test_json_preserves_nested_values_and_array_order(document):
    result = value_from(document, "files", "reader_options", (
        '{"options":{"names":["NA","",null],"enabled":false,"n":0},'
        '"skiprows":[2,0],"unicode":"測定","large":9223372036854775807}'
    ))
    assert result == {
        "options": {"names": ["NA", "", None], "enabled": False, "n": 0},
        "skiprows": [2, 0], "unicode": "測定", "large": 2**63 - 1,
    }


@pytest.mark.parametrize("raw", [
    "[]", "null", "false", "{bad}", '{"a":1,"a":2}', '{"x":{"a":1,"a":2}}',
    '{"a":NaN}', '{"a":Infinity}', '{"a":1e309}', '{"a":1e-999}',
    '{"a":"\\ud800"}', {}, None,
])
def test_invalid_json_objects_are_not_silently_normalized(document, raw):
    with pytest.raises(WorkbookParseError, match="reader_options"):
        value_from(document, "files", "reader_options", raw)


@pytest.mark.parametrize(("sheet", "column", "raw"), [
    ("sites", "row_role", "EDIT"), ("files", "reader_type", "xlsx"),
])
def test_control_enums_are_checked(document, sheet, column, raw):
    with pytest.raises(WorkbookParseError, match=column):
        value_from(document, sheet, column, raw)


@pytest.mark.parametrize(("kind", "raw"), [("f", "=1+1"), ("f", None), ("e", "#REF!"), ("d", START)])
def test_excel_special_cells_are_rejected_even_when_nullable(document, kind, raw):
    with pytest.raises(WorkbookParseError, match="description"):
        value_from(document, "sites", "description", raw, kind)


def test_errors_are_collected_in_schema_order_with_original_cells(document):
    document = with_cell(document, "sites", "name", 12)
    document = with_cell(document, "variables", "derived", "yes")
    with pytest.raises(WorkbookParseError) as caught:
        parse_workbook(document)
    assert [(issue.sheet, issue.coordinate) for issue in caught.value.issues] == [
        ("sites", "E7"), ("variables", "E7"),
    ]


def test_frames_do_not_coerce_mixed_null_and_large_id_to_float(document):
    document = with_cell(document, "sites", "site_id", str(2**63 - 1))
    first = document.sheets["sites"][0]
    second = {key: replace(cell, coordinate=cell.coordinate.replace("7", "9")) for key, cell in first.items()}
    second["site_id"] = replace(second["site_id"], value=None, data_type="n")
    document.sheets["sites"] = (first, second)
    result = parse_workbook(document)
    assert result.frames["sites"]["site_id"].tolist() == [2**63 - 1, None]
    assert type(result.frames["sites"].at[0, "site_id"]) is int
    assert result.coordinates["sites"].at[1, "site_id"] == "A9"


def test_input_is_unchanged_and_outputs_do_not_share_json_objects(document):
    document = with_cell(document, "files", "reader_options", '{"skiprows":[0,2]}')
    before = deepcopy(document)
    first = parse_workbook(document)
    second = parse_workbook(document)
    first.frames["files"].at[0, "reader_options"]["skiprows"].append(3)
    first.coordinates["files"].at[0, "path"] = "changed"
    assert second.frames["files"].at[0, "reader_options"] == {"skiprows": [0, 2]}
    assert document == before


def test_parsing_does_not_resolve_relationships_or_validate_scope(document):
    document = with_cell(document, "sensors", "sensor_model", " missing_model ")
    result = parse_workbook(document)
    assert result.frames["sensors"].at[0, "sensor_model"] == "missing_model"
    assert result.frames["sensor_models"].empty


def test_parser_does_not_open_files_or_databases(document, monkeypatch):
    def fail(*args, **kwargs):
        pytest.fail("parser attempted external IO")

    monkeypatch.setattr("builtins.open", fail)
    monkeypatch.setattr("dendroflow.configuration.workbook.source.connect", fail)
    monkeypatch.setattr("dendroflow.database.connect", fail)
    assert len(parse_workbook(document).frames) == 11
