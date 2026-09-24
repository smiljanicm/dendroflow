from dataclasses import replace
from datetime import datetime, timezone
from uuid import UUID

import pandas as pd
import pytest
from openpyxl.utils import get_column_letter
from pandas.testing import assert_frame_equal

from dendroflow.configuration.workbook.parsing import parse_workbook
from dendroflow.configuration.workbook.preparation import prepare_workbook_export
from dendroflow.configuration.workbook.reader import (
    WorkbookCell,
    WorkbookDocument,
    WorkbookMetadata,
    read_workbook,
)
from dendroflow.configuration.workbook.schema import RESOURCE_SHEETS, get_sheet_spec
from dendroflow.configuration.workbook.source import SOURCE_COLUMNS
from dendroflow.configuration.workbook.validation import (
    WorkbookValidationError,
    validate_workbook,
)
from dendroflow.configuration.workbook.writer import write_workbook_export

START = datetime(2025, 1, 1, tzinfo=timezone.utc)


@pytest.fixture
def source():
    records = {
        "sites": [(1, "region", "Region", None, None, None, None),
                  (2, "north", "North", None, None, None, 1),
                  (3, "south", "South", None, None, None, 1)],
        "location_types": [(10, "well", None)],
        "sensor_types": [(20, "logger", None)],
        "variables": [(30, "level", False, None)],
        "sensor_models": [(40, "Maker", "Model", 20)],
        "sensors": [(50, "000001", 40, None)],
        "locations": [(60, 2, 10, None, None, None, None), (61, 3, 10, None, None, None, None)],
        "location_labels": [(70, 60, "North", START, None), (71, 61, "South", START, None)],
        "deployments": [(80, 50, 60, 30, START, None), (81, 50, 61, 30, START, None)],
        "files": [(90, "/data/shared.csv", "UTC", "%Y-%m-%d", {"reader": "csv", "options": {}})],
        "interfaces": [(100, 90, 80, "time", "north", "cm"), (101, 90, 81, "time", "south", "cm")],
    }
    return {name: pd.DataFrame(records[name], columns=columns, dtype=object)
            for name, columns in SOURCE_COLUMNS.items()}


def parsed_export(source, site_ids=None):
    export = prepare_workbook_export(source, site_ids=site_ids)
    sheets = {}
    for spec in RESOURCE_SHEETS:
        rows = []
        for index, record in enumerate(export.frames[spec.name].to_dict("records"), start=2):
            rows.append({name: WorkbookCell(
                spec.name, f"{get_column_letter(column)}{index}", value,
                "b" if type(value) is bool else "n" if value is None or type(value) in {int, float} else "s",
            ) for column, (name, value) in enumerate(record.items(), start=1)})
        sheets[spec.name] = tuple(rows)
    metadata = WorkbookMetadata(1, UUID(int=1), "local-dev", START, export.scope, export.site_ids)
    return parse_workbook(WorkbookDocument(metadata, sheets))


@pytest.fixture
def parsed(source):
    return parsed_export(source)


@pytest.fixture
def scoped(source):
    return parsed_export(source, [2])


def add_row(parsed, sheet, row):
    index = len(parsed.frames[sheet])
    spec = get_sheet_spec(sheet)
    parsed.frames[sheet] = pd.concat([
        parsed.frames[sheet], pd.DataFrame([row], columns=spec.headers, dtype=object),
    ], ignore_index=True)
    coordinates = {name: f"{get_column_letter(column)}{index + 100}"
                   for column, name in enumerate(spec.headers, start=1)}
    parsed.coordinates[sheet] = pd.concat([
        parsed.coordinates[sheet], pd.DataFrame([coordinates], dtype=object),
    ], ignore_index=True)


def remove_rows(parsed, sheet, indexes):
    parsed.frames[sheet] = parsed.frames[sheet].drop(index=indexes).reset_index(drop=True)
    parsed.coordinates[sheet] = parsed.coordinates[sheet].drop(index=indexes).reset_index(drop=True)


def rejects(parsed, message):
    with pytest.raises(WorkbookValidationError, match=message) as caught:
        validate_workbook(parsed)
    return caught.value.issues


@pytest.mark.parametrize("selection", [None, [2], [1, 2, 3]])
def test_saved_export_passes_full_workbook_pipeline(source, tmp_path, selection):
    export = prepare_workbook_export(source, site_ids=selection)
    path = write_workbook_export(tmp_path / "control.xlsx", export, target_environment="local-dev")
    parsed = parse_workbook(read_workbook(path, expected_environment="local-dev"))
    assert validate_workbook(parsed) is None


def test_empty_template_and_new_related_setup_are_valid(parsed):
    new = replace(parsed, metadata=replace(parsed.metadata, scope="template", exported_at=None))
    for spec in RESOURCE_SHEETS:
        new.frames[spec.name][spec.id_column] = pd.Series([None] * len(new.frames[spec.name]), dtype=object)
    assert validate_workbook(new) is None
    for spec in RESOURCE_SHEETS:
        remove_rows(new, spec.name, list(new.frames[spec.name].index))
    assert validate_workbook(new) is None


@pytest.mark.parametrize("sheet", [spec.name for spec in RESOURCE_SHEETS])
@pytest.mark.parametrize("field", ["id", "ref"])
def test_duplicate_identities_are_located_on_every_conflicting_row(parsed, sheet, field):
    spec = get_sheet_spec(sheet)
    first = parsed.frames[sheet].iloc[0].to_dict()
    duplicate = dict(first)
    if field == "id":
        duplicate["ref"] = "another_ref"
    else:
        duplicate[spec.id_column] = 999
    add_row(parsed, sheet, duplicate)
    issues = rejects(parsed, "duplicate value")
    name = spec.id_column if field == "id" else "ref"
    coordinates = {issue.coordinate for issue in issues if issue.sheet == sheet and issue.message.startswith(name + ":")}
    assert parsed.coordinates[sheet].iloc[0][name] in coordinates
    assert parsed.coordinates[sheet].iloc[-1][name] in coordinates


REFERENCES = [(spec.name, column.name) for spec in RESOURCE_SHEETS for column in spec.fields if column.target_sheet]


@pytest.mark.parametrize(("sheet", "field"), REFERENCES)
def test_every_relationship_requires_a_reference_in_its_target_sheet(parsed, sheet, field):
    parsed.frames[sheet].at[0, field] = "missing"
    issues = rejects(parsed, "unresolved reference")
    assert any(issue.sheet == sheet and issue.coordinate == parsed.coordinates[sheet].at[0, field]
               for issue in issues)


def test_ambiguous_reference_is_not_silently_bound_to_first_match(parsed):
    row = parsed.frames["sensor_models"].iloc[0].to_dict()
    row["sensor_model_id"] = 999
    add_row(parsed, "sensor_models", row)
    issues = rejects(parsed, "ambiguous reference")
    assert any(issue.sheet == "sensors" and issue.coordinate == "E2" for issue in issues)


def test_aliases_are_case_sensitive_but_only_unique_within_each_sheet(parsed):
    parsed.frames["variables"].at[0, "ref"] = "sensor_50"
    parsed.frames["deployments"]["variable"] = "sensor_50"
    assert validate_workbook(parsed) is None
    parsed.frames["deployments"].at[0, "variable"] = "Sensor_50"
    rejects(parsed, "unresolved reference to variables")


@pytest.mark.parametrize("parent", ["site_1", "site_2"])
def test_site_parent_cycles_are_rejected(parsed, parent):
    parsed.frames["sites"].at[0, "parent"] = parent
    issues = rejects(parsed, "ancestry contains a cycle")
    assert any(issue.sheet == "sites" and issue.coordinate == "I2" for issue in issues)


def test_missing_unused_rows_do_not_request_deletion(parsed):
    remove_rows(parsed, "location_labels", [0, 1])
    assert validate_workbook(parsed) is None  # Existing histories are verified against DB in G4.


def test_all_scope_rows_cannot_be_changed_to_reference(parsed):
    parsed.frames["sites"].at[0, "row_role"] = "reference"
    rejects(parsed, "expected edit")


def test_template_cannot_contain_existing_ids(parsed):
    template = replace(parsed, metadata=replace(parsed.metadata, scope="template", exported_at=None))
    rejects(template, "template rows must have blank database IDs")


@pytest.mark.parametrize("sheet", ["location_types", "sensor_types", "variables", "sensor_models", "sensors", "files", "sites"])
def test_scoped_existing_shared_and_ancestor_rows_stay_reference(scoped, sheet):
    scoped.frames[sheet].at[0, "row_role"] = "edit"
    rejects(scoped, "expected reference")


@pytest.mark.parametrize("sheet", ["locations", "location_labels", "deployments", "interfaces"])
def test_scoped_owned_rows_stay_edit(scoped, sheet):
    scoped.frames[sheet].at[0, "row_role"] = "reference"
    rejects(scoped, "expected edit")


def test_selected_site_row_is_required(scoped):
    remove_rows(scoped, "sites", [1])
    issues = rejects(scoped, "selected site 2 is missing")
    assert any(issue.sheet == "workbook_info" and issue.coordinate is None for issue in issues)


def test_new_sites_require_all_or_template_scope(scoped):
    scoped.frames["sites"].at[1, "site_id"] = None
    rejects(scoped, "new sites require all or template")


def test_scoped_relationship_cannot_move_to_ancestor_site(scoped):
    scoped.frames["locations"].at[0, "site"] = "site_1"
    rejects(scoped, "outside the selected sites")


def test_unrelated_shared_record_does_not_belong_in_site_scope(scoped):
    row = scoped.frames["variables"].iloc[0].to_dict()
    row.update(variable_id=999, ref="unrelated")
    add_row(scoped, "variables", row)
    rejects(scoped, "not a selected site or a required dependency")


def test_shared_file_does_not_broaden_scope_to_other_interfaces(scoped, parsed):
    for name in ("sites", "locations", "deployments", "interfaces"):
        row = parsed.frames[name].iloc[-1].to_dict()
        if name == "sites":
            row["row_role"] = "reference"
        add_row(scoped, name, row)
    rejects(scoped, "outside the selected sites")


def test_new_supporting_resources_are_allowed_when_connected_to_selected_site(scoped):
    scoped.frames["sensors"].at[0, "sensor_id"] = None
    scoped.frames["sensors"].at[0, "row_role"] = "edit"
    assert validate_workbook(scoped) is None
    scoped.frames["sensors"].at[0, "row_role"] = "reference"
    rejects(scoped, "expected edit")


def test_new_unattached_supporting_resource_requires_broader_scope(scoped):
    row = scoped.frames["variables"].iloc[0].to_dict()
    row.update(variable_id=None, ref="new_unused", row_role="edit")
    add_row(scoped, "variables", row)
    rejects(scoped, "not a selected site or a required dependency")


def test_new_location_with_one_new_initial_and_extra_label_is_valid(parsed):
    parsed.frames["locations"].at[0, "location_id"] = None
    parsed.frames["location_labels"].at[0, "location_label_id"] = None
    extra = parsed.frames["location_labels"].iloc[0].to_dict()
    extra.update(ref="later_label", is_initial=False)
    add_row(parsed, "location_labels", extra)
    assert validate_workbook(parsed) is None


@pytest.mark.parametrize("defect", ["missing", "not_initial", "multiple", "existing_initial"])
def test_new_location_needs_exactly_one_new_initial_label(parsed, defect):
    parsed.frames["locations"].at[0, "location_id"] = None
    parsed.frames["location_labels"].at[0, "location_label_id"] = None
    if defect == "missing":
        remove_rows(parsed, "location_labels", [0])
    elif defect == "not_initial":
        parsed.frames["location_labels"].at[0, "is_initial"] = False
    elif defect == "existing_initial":
        parsed.frames["location_labels"].at[0, "location_label_id"] = 70
    else:
        row = parsed.frames["location_labels"].iloc[0].to_dict()
        row["ref"] = "extra_initial"
        add_row(parsed, "location_labels", row)
    rejects(parsed, "exactly one new initial label")


def test_existing_location_accepts_additional_new_noninitial_label(parsed):
    row = parsed.frames["location_labels"].iloc[0].to_dict()
    row.update(location_label_id=None, ref="additional", is_initial=False)
    add_row(parsed, "location_labels", row)
    assert validate_workbook(parsed) is None
    parsed.frames["location_labels"].at[2, "is_initial"] = True
    rejects(parsed, "new label for an existing location must not be initial")


def test_existing_location_cannot_have_multiple_present_initial_markers(parsed):
    row = parsed.frames["location_labels"].iloc[0].to_dict()
    row.update(location_label_id=999, ref="duplicate_initial")
    add_row(parsed, "location_labels", row)
    rejects(parsed, "multiple initial labels")


def test_new_file_requires_new_interface_but_existing_file_can_be_unreferenced(parsed):
    parsed.frames["files"].at[0, "file_id"] = None
    rejects(parsed, "new file requires at least one new interface")
    parsed.frames["interfaces"].at[0, "interface_id"] = None
    assert validate_workbook(parsed) is None
    parsed.frames["files"].at[0, "file_id"] = 90
    remove_rows(parsed, "interfaces", [0, 1])
    assert validate_workbook(parsed) is None


def test_validation_preserves_frames_and_handles_reordering(parsed, monkeypatch):
    def fail(*args, **kwargs):
        pytest.fail("validation attempted external IO")

    monkeypatch.setattr("builtins.open", fail)
    monkeypatch.setattr("dendroflow.database.connect", fail)
    monkeypatch.setattr("dendroflow.configuration.workbook.source.connect", fail)
    for name in parsed.frames:
        parsed.frames[name] = parsed.frames[name].iloc[::-1, ::-1].reset_index(drop=True)
        parsed.coordinates[name] = parsed.coordinates[name].iloc[::-1, ::-1].reset_index(drop=True)
    before = {name: frame.copy(deep=True) for name, frame in parsed.frames.items()}
    coordinates = {name: frame.copy(deep=True) for name, frame in parsed.coordinates.items()}
    assert validate_workbook(parsed) is None
    for name, frame in before.items():
        assert_frame_equal(parsed.frames[name], frame)
        assert_frame_equal(parsed.coordinates[name], coordinates[name])


def test_changed_existing_values_are_not_claimed_as_verified_against_database(parsed):
    parsed.frames["sites"].at[1, "site_id"] = 999
    parsed.frames["location_labels"].at[0, "label"] = "Changed"
    parsed.frames["location_labels"].at[0, "is_initial"] = False
    assert validate_workbook(parsed) is None  # G4 must reject stale IDs and unsupported edits.
