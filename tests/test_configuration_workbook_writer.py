from uuid import UUID

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

from dendroflow.configuration.workbook.schema import (
    FORMAT_VERSION,
    GUIDE_SHEET,
    INFO_KEYS,
    INFO_SHEET,
    RESOURCE_SHEETS,
    SHEET_NAMES,
    CellKind,
    collect_header_issues,
)
from dendroflow.configuration.workbook.writer import write_workbook_template


def read_info(workbook):
    return dict(workbook[INFO_SHEET].iter_rows(min_row=2, values_only=True))


def test_saved_template_matches_contract_and_has_no_resource_records(tmp_path):
    path = tmp_path / "control.xlsx"
    assert write_workbook_template(path, target_environment="local-dev") == path
    workbook = load_workbook(path)
    try:
        assert tuple(workbook.sheetnames) == SHEET_NAMES
        headers = {
            sheet.title: tuple(cell.value for cell in sheet[1])
            for sheet in workbook if sheet.title != GUIDE_SHEET
        }
        assert collect_header_issues(headers) == ()
        for spec in RESOURCE_SHEETS:
            assert not any(
                value is not None
                for row in workbook[spec.name].iter_rows(min_row=2, values_only=True)
                for value in row
            )
    finally:
        workbook.close()


def test_template_metadata_identifies_environment_without_claiming_export(tmp_path):
    path = write_workbook_template(
        tmp_path / "control.xlsx", target_environment=" local-dev ",
    )
    workbook = load_workbook(path)
    try:
        info = read_info(workbook)
        assert tuple(info) == INFO_KEYS
        assert info["format_version"] == FORMAT_VERSION
        assert UUID(info["workbook_id"]).version == 4
        assert info["target_environment"] == "local-dev"
        assert info["exported_at"] is None
        assert info["scope"] == "template"
        assert info["site_ids"] == "[]"
    finally:
        workbook.close()


def test_each_template_has_its_own_workbook_id(tmp_path):
    ids = []
    for name in ("first.xlsx", "second.xlsx"):
        path = write_workbook_template(tmp_path / name, target_environment="local-dev")
        workbook = load_workbook(path)
        try:
            ids.append(read_info(workbook)["workbook_id"])
        finally:
            workbook.close()
    assert ids[0] != ids[1]


@pytest.mark.parametrize("environment", ["=1+1", "+SUM(A1)", "@local", "#REF!"])
def test_environment_is_saved_as_literal_text(environment, tmp_path):
    path = write_workbook_template(tmp_path / "control.xlsx", target_environment=environment)
    workbook = load_workbook(path, data_only=False)
    try:
        cells = {
            row[0].value: row[1]
            for row in workbook[INFO_SHEET].iter_rows(min_row=2)
        }
        assert cells["target_environment"].value == environment
        assert cells["target_environment"].data_type == "s"
        assert all(
            cell.data_type not in {"f", "e"}
            for sheet in workbook for row in sheet for cell in row
        )
    finally:
        workbook.close()


def test_text_formats_preserve_ids_references_and_timestamp_entry(tmp_path):
    path = write_workbook_template(tmp_path / "control.xlsx", target_environment="local-dev")
    workbook = load_workbook(path)
    try:
        for spec in RESOURCE_SHEETS:
            sheet = workbook[spec.name]
            assert sheet.freeze_panes == "A2"
            for index, column in enumerate(spec.columns, start=1):
                if column.kind in {
                    CellKind.ID, CellKind.TEXT, CellKind.TIMESTAMP,
                    CellKind.JSON, CellKind.REFERENCE,
                }:
                    assert sheet.cell(2, index).number_format == "@"
                    letter = get_column_letter(index)
                    assert sheet.column_dimensions[letter].number_format == "@"
    finally:
        workbook.close()


def test_dropdowns_survive_save_and_cover_new_rows(tmp_path):
    path = write_workbook_template(tmp_path / "control.xlsx", target_environment="local-dev")
    workbook = load_workbook(path)
    try:
        for sheet_name, column, choices in (
            ("sites", "C", '"edit,reference"'),
            ("variables", "E", '"TRUE,FALSE"'),
            ("location_labels", "H", '"TRUE,FALSE"'),
            ("files", "G", '"csv"'),
        ):
            validators = workbook[sheet_name].data_validations.dataValidation
            matching = [v for v in validators if v.formula1 == choices]
            assert len(matching) == 1
            assert f"{column}2" in matching[0]
            assert f"{column}1048576" in matching[0]
            assert matching[0].showDropDown is False
            assert matching[0].showErrorMessage is True
    finally:
        workbook.close()


def test_existing_file_is_never_overwritten(tmp_path):
    path = tmp_path / "control.xlsx"
    original = b"existing workbook contents"
    path.write_bytes(original)
    with pytest.raises(FileExistsError):
        write_workbook_template(path, target_environment="local-dev")
    assert path.read_bytes() == original


@pytest.mark.parametrize("name", ["control", "control.csv", "control.xls", "control.xlsm"])
def test_only_xlsx_output_is_supported(name, tmp_path):
    path = tmp_path / name
    with pytest.raises(ValueError, match="must end with .xlsx"):
        write_workbook_template(path, target_environment="local-dev")
    assert not path.exists()


def test_string_path_and_uppercase_extension_are_supported(tmp_path):
    path = tmp_path / "control.XLSX"
    assert write_workbook_template(str(path), target_environment="local-dev") == path
    assert path.is_file()


@pytest.mark.parametrize("environment", ["", "   ", "a" * 129, "local\ndev", "local\x00dev"])
def test_invalid_environment_does_not_create_output(environment, tmp_path):
    path = tmp_path / "control.xlsx"
    with pytest.raises(ValueError, match="target_environment"):
        write_workbook_template(path, target_environment=environment)
    assert not path.exists()


@pytest.mark.parametrize("environment", [None, 42])
def test_environment_requires_text(environment, tmp_path):
    path = tmp_path / "control.xlsx"
    with pytest.raises(TypeError, match="must be a string"):
        write_workbook_template(path, target_environment=environment)
    assert not path.exists()


def test_missing_output_directory_is_not_created_implicitly(tmp_path):
    path = tmp_path / "missing" / "control.xlsx"
    with pytest.raises(FileNotFoundError):
        write_workbook_template(path, target_environment="local-dev")
    assert not path.parent.exists()


def test_serialization_error_leaves_no_output(tmp_path, monkeypatch):
    path = tmp_path / "control.xlsx"

    def fail_save(self, destination):
        destination.write(b"partial serialization")
        raise OSError("serialization failed")

    monkeypatch.setattr(Workbook, "save", fail_save)
    with pytest.raises(OSError, match="serialization failed"):
        write_workbook_template(path, target_environment="local-dev")
    assert not path.exists()


def test_template_creation_does_not_connect_to_database(tmp_path, monkeypatch):
    def fail_connect(*args, **kwargs):
        pytest.fail("template generation attempted database access")

    monkeypatch.setattr("dendroflow.database.connect", fail_connect)
    path = write_workbook_template(tmp_path / "control.xlsx", target_environment="local-dev")
    assert path.is_file()
