import os
import subprocess
import sys
from dataclasses import dataclass
from pathlib import Path
from uuid import uuid4

import psycopg
import pytest
import yaml
from openpyxl import load_workbook
from psycopg import sql
from psycopg.types.json import Jsonb

from dendroflow.cli import configuration
from dendroflow.cli.main import main
from dendroflow.database import connect

pytestmark = [
    pytest.mark.integration,
    pytest.mark.skipif(
        os.getenv("DENDROFLOW_INTEGRATION") != "1",
        reason="Set DENDROFLOW_INTEGRATION=1 to run PostgreSQL integration tests",
    ),
]

# Fixed test SQL, ordered for deletion; all values are parameters.
METADATA_TABLES = (
    ("deployments", "deployment_id",
     "variable_id IN (SELECT variable_id FROM variables WHERE variable = ANY(%s))"),
    ("sensors", "sensor_id", "serial_number = ANY(%s)"),
    ("sensor_models", "sensor_model_id", "model = ANY(%s)"),
    ("sensor_types", "sensor_type_id", "type = ANY(%s)"),
    ("variables", "variable_id", "variable = ANY(%s)"),
    ("location_labels", "location_label_id",
     (
         "location_id IN (SELECT location_id FROM locations WHERE site_id IN "
         "(SELECT site_id FROM sites WHERE site_code = ANY(%s)))"
     )),
    ("locations", "location_id",
     "site_id IN (SELECT site_id FROM sites WHERE site_code = ANY(%s))"),
    ("location_types", "location_type_id", "type = ANY(%s)"),
    ("sites", "site_id", "site_code = ANY(%s)"),
)


@dataclass
class CliCase:
    tag: str
    path: Path
    source: Path

    @property
    def codes(self):
        return [self.tag, f"{self.tag}_new", f"{self.tag}_corrected"]


@pytest.fixture
def cli_case(tmp_path):
    case = CliCase(
        f"cli_config_{uuid4().hex}",
        tmp_path / "config.yaml",
        tmp_path / "source.csv",
    )
    try:
        yield case
    finally:
        with connect("dendroflow_raw") as connection:
            connection.execute(
                """
                DELETE FROM sensor_file_interfaces
                WHERE file_id IN (SELECT file_id FROM files WHERE filepath = %s)
                """,
                (str(case.source),),
            )
            connection.execute(
                "DELETE FROM files WHERE filepath = %s",
                (str(case.source),),
            )
        with connect("dendroflow_metadata") as connection:
            for table, _, condition in METADATA_TABLES:
                connection.execute(
                    sql.SQL("DELETE FROM {} WHERE {}").format(
                        sql.Identifier(table), sql.SQL(condition),
                    ),
                    (case.codes,),
                )


def write_config(case, document):
    case.path.write_text(yaml.safe_dump(document, sort_keys=False), encoding="utf-8")


def full_config(case):
    tag = case.tag
    return {
        "sites": [{"ref": "site", "site_code": tag, "name": tag}],
        "location_types": [{"ref": "place", "type": tag}],
        "sensor_types": [{"ref": "kind", "type": tag}],
        "variables": [{"ref": "variable", "variable": tag}],
        "sensor_models": [{
            "ref": "model", "manufacturer": "CLI integration",
            "model": tag, "sensor_type": "kind",
        }],
        "sensors": [{
            "ref": "sensor", "serial_number": tag, "sensor_model": "model",
        }],
        "locations": [{
            "ref": "location", "site": "site", "location_type": "place",
            "initial_label": {
                "label": "L1", "valid_from": "2026-01-01T00:00:00Z",
            },
        }],
        "deployments": [{
            "ref": "deployment", "sensor": "sensor",
            "location": "location", "variable": "variable",
            "valid_from": "2026-01-01T00:00:00Z",
        }],
        "files": [{
            "path": str(case.source),
            "timestamp": {"timezone": "UTC", "format": "%Y-%m-%d %H:%M:%S"},
            "reader": {"type": "csv", "options": {}},
            "interfaces": [{
                "deployment": "deployment", "timestamp_column": "TIMESTAMP",
                "values_column": "value", "unit": "cm",
            }],
        }],
    }


def metadata_state(case):
    with connect("dendroflow_metadata") as connection:
        return {
            table: tuple(
                row[0] for row in connection.execute(
                    sql.SQL("SELECT {} FROM {} WHERE {} ORDER BY {}").format(
                        sql.Identifier(primary_key), sql.Identifier(table),
                        sql.SQL(condition), sql.Identifier(primary_key),
                    ),
                    (case.codes,),
                ).fetchall()
            )
            for table, primary_key, condition in METADATA_TABLES
        }


def raw_state(case):
    with connect("dendroflow_raw") as connection:
        files = connection.execute(
            "SELECT file_id FROM files WHERE filepath = %s ORDER BY file_id",
            (str(case.source),),
        ).fetchall()
        interfaces = connection.execute(
            """
            SELECT interface_id, file_id, deployment_id
            FROM sensor_file_interfaces
            WHERE file_id IN (SELECT file_id FROM files WHERE filepath = %s)
            ORDER BY interface_id
            """,
            (str(case.source),),
        ).fetchall()
    return files, interfaces


def workbook_mutation_state(case):
    """Capture fixture rows and xmin values to detect inserts/updates/deletes."""
    with connect("dendroflow_metadata") as connection:
        metadata = {
            table: tuple(
                row
                for row in connection.execute(
                    sql.SQL("SELECT {}, xmin::text FROM {} WHERE {} ORDER BY {}").format(
                        sql.Identifier(primary_key),
                        sql.Identifier(table),
                        sql.SQL(condition),
                        sql.Identifier(primary_key),
                    ),
                    (case.codes,),
                ).fetchall()
            )
            for table, primary_key, condition in METADATA_TABLES
        }
    with connect("dendroflow_raw") as connection:
        files = connection.execute(
            "SELECT file_id, xmin::text FROM files WHERE filepath = %s ORDER BY file_id",
            (str(case.source),),
        ).fetchall()
        interfaces = connection.execute(
            """
            SELECT interface_id, xmin::text, file_id, deployment_id
            FROM sensor_file_interfaces
            WHERE file_id IN (SELECT file_id FROM files WHERE filepath = %s)
            ORDER BY interface_id
            """,
            (str(case.source),),
        ).fetchall()
    return metadata, tuple(files), tuple(interfaces)


def assert_empty(case):
    assert not any(metadata_state(case).values())
    assert raw_state(case) == ([], [])


def seed_site(case):
    with connect("dendroflow_metadata") as connection:
        return connection.execute(
            "INSERT INTO sites (site_code, name) VALUES (%s, %s) RETURNING site_id",
            (case.tag, "Original"),
        ).fetchone()[0]


def site_state(case):
    with connect("dendroflow_metadata") as connection:
        rows = connection.execute(
            "SELECT site_code, site_id, name FROM sites WHERE site_code = ANY(%s)",
            (case.codes,),
        ).fetchall()
    return {code: (database_id, name) for code, database_id, name in rows}


def run_cli(*args):
    return subprocess.run(
        [sys.executable, "-m", "dendroflow", *map(str, args)],
        capture_output=True,
        text=True,
        check=False,
        timeout=30,
    )


def assert_success(result):
    assert result.returncode == 0, result.stdout + result.stderr
    assert result.stderr == ""


def test_yaml_cli_workflow_and_repeat_apply(cli_case):
    case = cli_case
    write_config(case, full_config(case))

    assert_success(run_cli("config", "validate", case.path))
    preview = run_cli("config", "plan", case.path)
    assert_success(preview)
    assert "Preparation checks passed" in preview.stdout
    assert_empty(case)

    applied = run_cli("config", "apply", case.path, "--yes")
    assert_success(applied)
    assert "METADATA: COMMITTED" in applied.stdout
    assert "RAW: COMMITTED" in applied.stdout

    metadata_before = metadata_state(case)
    raw_before = raw_state(case)
    assert all(len(ids) == 1 for ids in metadata_before.values())
    files, interfaces = raw_before
    assert len(files) == len(interfaces) == 1
    assert interfaces[0][1] == files[0][0]
    assert interfaces[0][2] == metadata_before["deployments"][0]
    assert not case.source.exists()

    repeated_plan = run_cli("config", "plan", case.path)
    assert_success(repeated_plan)
    assert "Total: CREATE=0, REUSE=11, UPDATE=0" in repeated_plan.stdout

    # No --yes is needed because all operations now resolve to REUSE.
    repeated_apply = run_cli("config", "apply", case.path)
    assert_success(repeated_apply)
    assert "METADATA: NOT_REQUIRED" in repeated_apply.stdout
    assert "RAW: NOT_REQUIRED" in repeated_apply.stdout
    assert metadata_state(case) == metadata_before
    assert raw_state(case) == raw_before


def test_workbook_unchanged_export_validate_convert_plan_apply_is_read_only(
    cli_case, monkeypatch, capsys,
):
    case = cli_case
    monkeypatch.setenv("DENDROFLOW_ENVIRONMENT", "local")
    write_config(case, full_config(case))
    seeded = run_cli("config", "apply", case.path, "--yes")
    assert_success(seeded)

    site_id = metadata_state(case)["sites"][0]
    workbook_path = case.path.with_suffix(".xlsx")
    yaml_path = case.path.with_name("workbook-config.yaml")
    before = workbook_mutation_state(case)
    writes = []
    original_connect = psycopg.connect

    class ObservedConnection:
        def __init__(self, connection, database):
            self._connection = connection
            self._database = database

        def __enter__(self):
            self._connection.__enter__()
            return self

        def __exit__(self, *args):
            return self._connection.__exit__(*args)

        def __getattr__(self, name):
            return getattr(self._connection, name)

        def execute(self, query, *args, **kwargs):
            cursor = self._connection.execute(query, *args, **kwargs)
            status = cursor.statusmessage or ""
            if status.partition(" ")[0].upper() in {
                "INSERT", "UPDATE", "DELETE", "MERGE", "TRUNCATE",
            }:
                writes.append((self._database, status))
            return cursor

    def observed_connect(*args, **kwargs):
        connection = original_connect(*args, **kwargs)
        return ObservedConnection(connection, kwargs.get("dbname", "unknown"))

    monkeypatch.setattr(psycopg, "connect", observed_connect)

    commands = (
        ["config", "workbook", "export", "--site-id", str(site_id), str(workbook_path)],
        ["config", "workbook", "validate", str(workbook_path)],
        ["config", "workbook", "convert", str(workbook_path), str(yaml_path)],
        ["config", "plan", str(yaml_path)],
        ["config", "apply", str(yaml_path)],
    )
    for command in commands:
        result = main(command)
        captured = capsys.readouterr()
        assert result == 0, captured.out + captured.err

    assert "Workbook valid:" in captured.out or "METADATA: NOT_REQUIRED" in captured.out
    assert not writes
    assert workbook_mutation_state(case) == before


def test_workbook_supported_updates_and_new_related_rows(cli_case, monkeypatch, capsys):
    case = cli_case
    monkeypatch.setenv("DENDROFLOW_ENVIRONMENT", "local")
    document = full_config(case)
    document["sites"][0]["description"] = "Description to clear"
    write_config(case, document)
    seeded = run_cli("config", "apply", case.path, "--yes")
    assert_success(seeded)

    site_id = metadata_state(case)["sites"][0]
    before = workbook_mutation_state(case)
    workbook_path = case.path.with_suffix(".xlsx")
    yaml_path = case.path.with_name("workbook-changes.yaml")
    assert main([
        "config", "workbook", "export", "--site-id", str(site_id),
        str(workbook_path),
    ]) == 0
    capsys.readouterr()

    with connect("dendroflow_metadata") as connection:
        original_relationship = connection.execute(
            """
            SELECT d.sensor_id, d.location_id
            FROM deployments AS d
            JOIN variables AS v USING (variable_id)
            WHERE v.variable = %s
            """,
            (case.tag,),
        ).fetchone()

    workbook = load_workbook(workbook_path)
    try:
        sites = workbook["sites"]
        site_columns = {cell.value: cell.column for cell in sites[1]}
        site_row = next(
            row
            for row in range(2, sites.max_row + 1)
            if sites.cell(row, site_columns["site_id"]).value == str(site_id)
        )
        sites.cell(site_row, site_columns["site_code"], case.codes[2])
        sites.cell(site_row, site_columns["name"], "Revised site name")
        sites.cell(site_row, site_columns["description"]).value = None

        variables = workbook["variables"]
        variable_values = {
            "variable_id": None,
            "ref": "new_variable",
            "row_role": "edit",
            "variable": case.codes[1],
            "derived": False,
            "description": None,
        }
        variables.append([
            variable_values[cell.value]
            for cell in variables[1]
        ])

        deployments = workbook["deployments"]
        deployment_columns = {cell.value: cell.column for cell in deployments[1]}
        deployment_values = {
            "deployment_id": None,
            "ref": "new_deployment",
            "row_role": "edit",
            "sensor": deployments.cell(2, deployment_columns["sensor"]).value,
            "location": deployments.cell(2, deployment_columns["location"]).value,
            "variable": "new_variable",
            "valid_from": "2027-01-01T00:00:00Z",
            "valid_to": None,
        }
        deployments.append([
            deployment_values[cell.value]
            for cell in deployments[1]
        ])
        workbook.save(workbook_path)
    finally:
        workbook.close()

    assert main(["config", "workbook", "validate", str(workbook_path)]) == 0
    capsys.readouterr()
    converted = main([
        "config", "workbook", "convert", str(workbook_path), str(yaml_path),
    ])
    converted_output = capsys.readouterr()
    assert converted == 0, converted_output.out + converted_output.err

    planned = main(["config", "plan", str(yaml_path)])
    plan_output = capsys.readouterr().out
    assert planned == 0
    assert "CREATE=2" in plan_output
    assert "UPDATE=1" in plan_output

    unconfirmed = main(["config", "apply", str(yaml_path), "--yes"])
    unconfirmed_output = capsys.readouterr()
    assert unconfirmed == 3
    assert "--confirm-identity-changes is required" in unconfirmed_output.err
    assert workbook_mutation_state(case) == before

    applied = main([
        "config", "apply", str(yaml_path), "--yes",
        "--confirm-identity-changes",
    ])
    applied_output = capsys.readouterr()
    assert applied == 0, applied_output.out + applied_output.err
    assert "METADATA: COMMITTED" in applied_output.out

    assert site_state(case) == {
        case.codes[2]: (site_id, "Revised site name"),
    }
    with connect("dendroflow_metadata") as connection:
        description = connection.execute(
            "SELECT description FROM sites WHERE site_id = %s",
            (site_id,),
        ).fetchone()[0]
        variable = connection.execute(
            "SELECT variable_id FROM variables WHERE variable = %s",
            (case.codes[1],),
        ).fetchone()[0]
        deployment = connection.execute(
            """
            SELECT sensor_id, location_id
            FROM deployments
            WHERE variable_id = %s
            """,
            (variable,),
        ).fetchone()
    assert description is None
    assert deployment == original_relationship
    assert len(metadata_state(case)["deployments"]) == len(before[0]["deployments"]) + 1


def test_workbook_adds_interface_to_existing_file(cli_case, monkeypatch, capsys):
    case = cli_case
    monkeypatch.setenv("DENDROFLOW_ENVIRONMENT", "local")
    write_config(case, full_config(case))
    seeded = run_cli("config", "apply", case.path, "--yes")
    assert_success(seeded)

    site_id = metadata_state(case)["sites"][0]
    metadata_before, raw_before = metadata_state(case), raw_state(case)
    assert len(raw_before[0]) == len(raw_before[1]) == 1
    original_file_id = raw_before[0][0][0]

    workbook_path = case.path.with_suffix(".xlsx")
    yaml_path = case.path.with_name("workbook-interface.yaml")
    assert main([
        "config", "workbook", "export", "--site-id", str(site_id),
        str(workbook_path),
    ]) == 0
    capsys.readouterr()

    workbook = load_workbook(workbook_path)
    try:
        variables = workbook["variables"]
        variable_values = {
            "variable_id": None,
            "ref": "new_variable",
            "row_role": "edit",
            "variable": case.codes[1],
            "derived": False,
            "description": None,
        }
        variables.append([variable_values[cell.value] for cell in variables[1]])

        deployments = workbook["deployments"]
        deployment_columns = {cell.value: cell.column for cell in deployments[1]}
        deployment_values = {
            "deployment_id": None,
            "ref": "new_deployment",
            "row_role": "edit",
            "sensor": deployments.cell(2, deployment_columns["sensor"]).value,
            "location": deployments.cell(2, deployment_columns["location"]).value,
            "variable": "new_variable",
            "valid_from": "2027-01-01T00:00:00Z",
            "valid_to": None,
        }
        deployments.append([
            deployment_values[cell.value] for cell in deployments[1]
        ])

        interfaces = workbook["interfaces"]
        interface_columns = {cell.value: cell.column for cell in interfaces[1]}
        interface_values = {
            "interface_id": None,
            "ref": "new_interface",
            "row_role": "edit",
            "file": interfaces.cell(2, interface_columns["file"]).value,
            "deployment": "new_deployment",
            "timestamp_column": "TIMESTAMP",
            "values_column": "value_new",
            "unit": "cm",
        }
        interfaces.append([
            interface_values[cell.value] for cell in interfaces[1]
        ])
        workbook.save(workbook_path)
    finally:
        workbook.close()

    assert main(["config", "workbook", "validate", str(workbook_path)]) == 0
    capsys.readouterr()
    converted = main([
        "config", "workbook", "convert", str(workbook_path), str(yaml_path),
    ])
    converted_output = capsys.readouterr()
    assert converted == 0, converted_output.out + converted_output.err

    planned = main(["config", "plan", str(yaml_path)])
    plan_output = capsys.readouterr().out
    assert planned == 0
    assert "CREATE=2" in plan_output
    assert "UPDATE=0" in plan_output

    applied = main(["config", "apply", str(yaml_path), "--yes"])
    applied_output = capsys.readouterr()
    assert applied == 0, applied_output.out + applied_output.err
    assert "METADATA: COMMITTED" in applied_output.out
    assert "RAW: COMMITTED" in applied_output.out

    metadata_after = metadata_state(case)
    raw_after = raw_state(case)
    assert len(metadata_after["variables"]) == len(metadata_before["variables"]) + 1
    assert len(metadata_after["deployments"]) == len(metadata_before["deployments"]) + 1
    assert raw_after[0] == raw_before[0]
    assert len(raw_after[1]) == len(raw_before[1]) + 1
    with connect("dendroflow_metadata") as connection:
        new_deployment_id = connection.execute(
            """
            SELECT d.deployment_id
            FROM deployments AS d
            JOIN variables AS v USING (variable_id)
            WHERE v.variable = %s
            """,
            (case.codes[1],),
        ).fetchone()[0]
    added_interfaces = [row for row in raw_after[1] if row not in raw_before[1]]
    assert len(added_interfaces) == 1
    assert added_interfaces[0][1:] == (original_file_id, new_deployment_id)


def test_identity_confirmation_preserves_existing_id(cli_case, capsys):
    case = cli_case
    database_id = seed_site(case)
    corrected = f"{case.tag}_corrected"
    write_config(case, {
        "updates": {"sites": [{
            "update": {"site_code": case.tag},
            "set": {"site_code": corrected},
        }]},
    })

    assert main(["config", "apply", str(case.path), "--yes"]) == 3
    assert "--confirm-identity-changes is required" in capsys.readouterr().err
    assert site_state(case) == {case.tag: (database_id, "Original")}

    assert main([
        "config", "apply", str(case.path),
        "--yes", "--confirm-identity-changes",
    ]) == 0
    assert "METADATA: COMMITTED" in capsys.readouterr().out
    assert site_state(case) == {corrected: (database_id, "Original")}


def test_conflict_blocks_all_writes(cli_case, capsys):
    case = cli_case
    database_id = seed_site(case)
    write_config(case, {
        "sites": [
            {"site_code": f"{case.tag}_new", "name": "New"},
            {"site_code": case.tag, "name": "Conflicting declaration"},
        ],
    })

    assert main(["config", "apply", str(case.path), "--yes"]) == 2
    captured = capsys.readouterr()
    assert "CONFLICT" in captured.out
    assert "Apply blocked" in captured.err
    assert site_state(case) == {case.tag: (database_id, "Original")}
    assert raw_state(case) == ([], [])


def test_declined_apply_writes_nothing(cli_case, monkeypatch, capsys):
    case = cli_case
    write_config(case, full_config(case))
    monkeypatch.setattr("sys.stdin.isatty", lambda: True)
    monkeypatch.setattr("builtins.input", lambda prompt: "no")

    assert main(["config", "apply", str(case.path)]) == 3
    assert "Apply cancelled" in capsys.readouterr().err
    assert_empty(case)


def test_stale_update_rolls_back_earlier_create(cli_case, monkeypatch, capsys):
    case = cli_case
    database_id = seed_site(case)
    write_config(case, {
        "sites": [{"site_code": f"{case.tag}_new", "name": "New"}],
        "updates": {"sites": [{
            "update": {"site_code": case.tag}, "set": {"name": "Requested"},
        }]},
    })
    original_review = configuration.review_for_apply

    def change_after_review(plan, **kwargs):
        code = original_review(plan, **kwargs)
        assert code == 0
        with connect("dendroflow_metadata") as connection:
            connection.execute(
                "UPDATE sites SET name = %s WHERE site_id = %s",
                ("Concurrent", database_id),
            )
        return code

    monkeypatch.setattr(configuration, "review_for_apply", change_after_review)

    assert main(["config", "apply", str(case.path), "--yes"]) == 1
    captured = capsys.readouterr()
    assert "Database outcome: FAILED" in captured.out
    assert "STALE_PLAN" in captured.err
    assert site_state(case) == {case.tag: (database_id, "Concurrent")}


def test_raw_race_reports_partial_and_preserves_metadata(
    cli_case, monkeypatch, capsys,
):
    case = cli_case
    write_config(case, full_config(case))
    original_review = configuration.review_for_apply
    competing_ids = []

    def insert_after_review(plan, **kwargs):
        code = original_review(plan, **kwargs)
        assert code == 0
        with connect("dendroflow_raw") as connection:
            row = connection.execute(
                """
                INSERT INTO files (
                    filepath, timestamp_timezone, timestamp_format, reader_config
                )
                VALUES (%s, %s, %s, %s)
                RETURNING file_id
                """,
                (
                    str(case.source), "UTC", "%Y-%m-%d %H:%M:%S",
                    Jsonb({"reader": "csv", "options": {}}),
                ),
            ).fetchone()
            competing_ids.append(row[0])
        return code

    monkeypatch.setattr(configuration, "review_for_apply", insert_after_review)

    assert main(["config", "apply", str(case.path), "--yes"]) == 4
    captured = capsys.readouterr()
    outcome = captured.out.split("Database outcome:", 1)[1]
    assert outcome.startswith(" PARTIAL\n")
    assert "METADATA: COMMITTED" in outcome
    assert "RAW: FAILED" in outcome
    assert "CREATE file" not in outcome
    assert "CREATE interface" not in outcome
    assert "UniqueViolation" in captured.err
    assert all(len(ids) == 1 for ids in metadata_state(case).values())
    assert raw_state(case) == ([(competing_ids[0],)], [])
