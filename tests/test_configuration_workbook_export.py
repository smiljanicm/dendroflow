import json
from dataclasses import replace
from datetime import datetime, timedelta, timezone
from uuid import UUID

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from pandas.testing import assert_frame_equal

from dendroflow.configuration.workbook.preparation import prepare_workbook_export
from dendroflow.configuration.workbook.schema import (
    FORMAT_VERSION,
    INFO_KEYS,
    RESOURCE_SHEETS,
    SHEET_NAMES,
    CellKind,
)
from dendroflow.configuration.workbook.source import SOURCE_COLUMNS
from dendroflow.configuration.workbook.writer import write_workbook_export

START = datetime(2025, 1, 1, 2, 0, 0, 123456, tzinfo=timezone(timedelta(hours=2)))


def source_frames(records=None):
    records = records or {}
    return {
        name: pd.DataFrame(records.get(name, []), columns=columns, dtype=object)
        for name, columns in SOURCE_COLUMNS.items()
    }


@pytest.fixture
def frames():
    return source_frames({
        "sites": [
            (1, "region", "Region", None, None, None, None),
            (2, "000012", "Český les", "Line one\nLine two", 53.123456789012344, 0.0, 1),
        ],
        "location_types": [(10, "well", None)],
        "sensor_types": [(20, "logger", None)],
        "variables": [(30, "level", False, None)],
        "sensor_models": [(40, "Maker", "Model", 20)],
        "sensors": [(50, "000012", 40, None)],
        "locations": [(60, 2, 10, None, None, -0.0, 0)],
        "location_labels": [(70, 60, "Well", START, None)],
        "deployments": [(80, 50, 60, 30, START, None)],
        "files": [(90, "/data/測定.csv", "UTC", "%Y-%m-%d", {
            "reader": "csv", "options": {"skiprows": [0, 2], "na_values": ["NA", ""]},
        })],
        "interfaces": [(100, 90, 80, "time", "level", "cm")],
    })


@pytest.fixture
def export(frames):
    return prepare_workbook_export(frames)


@pytest.fixture
def save(tmp_path):
    opened = []

    def write(export, **kwargs):
        path = tmp_path / f"export-{len(opened)}.xlsx"
        write_workbook_export(path, export, target_environment="local-dev", **kwargs)
        workbook = load_workbook(path, data_only=False)
        opened.append(workbook)
        return workbook

    yield write
    for workbook in opened:
        workbook.close()


def info(workbook):
    return dict(workbook["workbook_info"].iter_rows(min_row=2, values_only=True))


def test_saved_export_preserves_all_prepared_cells_and_sheet_layout(export, save):
    workbook = save(export)
    assert tuple(workbook.sheetnames) == SHEET_NAMES
    for spec in RESOURCE_SHEETS:
        sheet = workbook[spec.name]
        assert tuple(cell.value for cell in sheet[1]) == spec.headers
        assert sheet.max_row == len(export.frames[spec.name]) + 1
        assert sheet.freeze_panes == "A2"
        assert sheet.auto_filter.ref == (
            f"A1:{get_column_letter(len(spec.headers))}{sheet.max_row}"
        )
        for row_index, values in enumerate(
            export.frames[spec.name].itertuples(index=False, name=None), start=2,
        ):
            for column_index, (column, value) in enumerate(zip(spec.columns, values), start=1):
                cell = sheet.cell(row_index, column_index)
                if value is None:
                    assert cell.value is None
                elif column.kind == CellKind.NUMBER:
                    assert cell.value == repr(value)
                    assert cell.data_type == "s"
                    assert float(cell.value) == value
                elif column.kind == CellKind.BOOLEAN:
                    assert cell.value is value
                    assert cell.data_type == "b"
                else:
                    assert cell.value == value
                    assert cell.data_type == "s"
                    assert cell.number_format == "@"
    assert workbook["deployments"]["G2"].value == "2025-01-01T00:00:00.123456Z"
    assert json.loads(workbook["files"]["H2"].value) == {
        "skiprows": [0, 2], "na_values": ["NA", ""],
    }


def test_site_export_retains_scope_and_reference_roles(frames, save):
    export = prepare_workbook_export(frames, site_ids=[2])
    workbook = save(export)
    assert info(workbook)["scope"] == "sites"
    assert json.loads(info(workbook)["site_ids"]) == ["2"]
    assert workbook["sites"]["C2"].value == "reference"
    assert workbook["sites"]["C3"].value == "edit"
    assert workbook["sites"]["I3"].value == "site_1"
    assert workbook["sites"]["A2"].fill.fgColor.rgb == "00E8EDF0"
    assert workbook["sites"]["A3"].fill.patternType is None
    assert workbook["files"]["C2"].value == "reference"
    assert workbook["interfaces"]["C2"].value == "edit"
    assert not workbook["sites"].protection.sheet


def test_export_metadata_has_unique_id_environment_and_utc_time(export, save):
    workbook = save(export, exported_at=START)
    values = info(workbook)
    assert tuple(values) == INFO_KEYS
    assert values["format_version"] == FORMAT_VERSION
    assert UUID(values["workbook_id"]).version == 4
    assert values["target_environment"] == "local-dev"
    assert values["scope"] == "all"
    assert values["site_ids"] == "[]"
    assert values["exported_at"] == "2025-01-01T00:00:00.123456Z"
    before = datetime.now(timezone.utc)
    second = info(save(export))
    after = datetime.now(timezone.utc)
    assert second["workbook_id"] != values["workbook_id"]
    assert before <= datetime.fromisoformat(second["exported_at"].replace("Z", "+00:00")) <= after


def test_empty_database_produces_all_headers_with_export_metadata(save):
    workbook = save(prepare_workbook_export(source_frames()))
    assert info(workbook)["scope"] == "all"
    assert info(workbook)["exported_at"]
    for spec in RESOURCE_SHEETS:
        sheet = workbook[spec.name]
        assert not any(
            value is not None for row in sheet.iter_rows(min_row=2, values_only=True)
            for value in row
        )
        assert sheet.auto_filter.ref == f"A1:{get_column_letter(len(spec.headers))}1"


def test_large_ids_leading_zeros_and_float_precision_survive_saved_file(save):
    key = 2**63 - 1
    export = prepare_workbook_export(source_frames({
        "sites": [(key, "000001", "NA", None, 53.123456789012344, 5e-324, None)],
    }), site_ids=[key])
    workbook = save(export)
    assert workbook["sites"]["A2"].value == str(key)
    assert workbook["sites"]["D2"].value == "000001"
    assert workbook["sites"]["E2"].value == "NA"
    assert workbook["sites"]["G2"].value == "53.123456789012344"
    assert workbook["sites"]["H2"].value == "5e-324"
    assert json.loads(info(workbook)["site_ids"]) == [str(key)]


@pytest.mark.parametrize("text", ["=1+1", "+SUM(A1)", "-1+2", "@SUM(A1)", "#REF!"])
def test_formula_and_error_looking_strings_are_literal(export, save, text):
    export.frames["sites"].at[0, "name"] = text
    workbook = save(export)
    assert workbook["sites"]["E2"].value == text
    assert workbook["sites"]["E2"].data_type == "s"
    assert all(
        cell.data_type not in {"f", "e"}
        for sheet in workbook for row in sheet for cell in row
    )


def test_writer_orders_headers_without_mutating_frames(export, save):
    for name, frame in export.frames.items():
        export.frames[name] = frame.iloc[:, ::-1].copy()
    before = {name: frame.copy(deep=True) for name, frame in export.frames.items()}
    workbook = save(export)
    for spec in RESOURCE_SHEETS:
        assert_frame_equal(export.frames[spec.name], before[spec.name])
        assert tuple(cell.value for cell in workbook[spec.name][1]) == spec.headers
    assert workbook["sites"]["E2"].value == "Region"


def test_guide_and_dropdowns_support_export_editing(export, save):
    workbook = save(export)
    guide = " ".join(str(row[0]) for row in workbook["guide"].values)
    assert "Retain existing IDs" in guide
    assert "decimal text" in guide
    assert "fresh export" in guide
    assert "This empty template" not in guide
    for name, column, choices in (
        ("sites", "C", '"edit,reference"'),
        ("variables", "E", '"TRUE,FALSE"'),
        ("files", "G", '"csv"'),
    ):
        validators = workbook[name].data_validations.dataValidation
        validation = next(item for item in validators if item.formula1 == choices)
        assert f"{column}2" in validation
        assert f"{column}1048576" in validation


@pytest.mark.parametrize("timestamp", [START.replace(tzinfo=None), "2025-01-01Z", 0])
def test_invalid_export_time_creates_no_file(export, tmp_path, timestamp):
    path = tmp_path / "control.xlsx"
    with pytest.raises(ValueError, match="timezone-aware"):
        write_workbook_export(path, export, target_environment="local", exported_at=timestamp)
    assert not path.exists()


@pytest.mark.parametrize(("scope", "site_ids"), [
    ("template", ()), ("all", (2,)), ("sites", ()), ("sites", (True,)),
    ("sites", (2, 1)), ("sites", (2, 2)), ("sites", (999,)),
])
def test_inconsistent_scope_creates_no_file(export, tmp_path, scope, site_ids):
    path = tmp_path / "control.xlsx"
    export = replace(export, scope=scope, site_ids=site_ids)
    with pytest.raises(ValueError):
        write_workbook_export(path, export, target_environment="local")
    assert not path.exists()


@pytest.mark.parametrize("defect", [
    "missing_sheet", "extra_sheet", "missing_header", "duplicate_header",
    "duplicate_id", "duplicate_ref", "invalid_id", "invalid_role", "reference_in_all",
])
def test_mutated_export_structure_is_rejected(export, tmp_path, defect):
    frame = export.frames["sites"]
    if defect == "missing_sheet":
        del export.frames["files"]
    elif defect == "extra_sheet":
        export.frames["unexpected"] = frame
    elif defect == "missing_header":
        export.frames["sites"] = frame.drop(columns="name")
    elif defect == "duplicate_header":
        export.frames["sites"] = pd.concat([frame, frame[["name"]]], axis=1)
    elif defect == "duplicate_id":
        frame.at[1, "site_id"] = "1"
    elif defect == "duplicate_ref":
        frame.at[1, "ref"] = frame.at[0, "ref"]
    elif defect == "invalid_id":
        frame.at[0, "site_id"] = "01"
    elif defect == "invalid_role":
        frame.at[0, "row_role"] = "other"
    else:
        frame.at[0, "row_role"] = "reference"
    path = tmp_path / "control.xlsx"
    with pytest.raises(ValueError):
        write_workbook_export(path, export, target_environment="local")
    assert not path.exists()


@pytest.mark.parametrize(("sheet", "column", "value"), [
    ("sites", "name", None), ("sites", "name", " padded "),
    ("sites", "name", "x" * 32768), ("sites", "latitude", float("nan")),
    ("sites", "latitude", True), ("variables", "derived", "FALSE"),
])
def test_invalid_mutated_cell_is_reported_before_file_creation(
    export, tmp_path, sheet, column, value,
):
    export.frames[sheet].at[0, column] = value
    path = tmp_path / "control.xlsx"
    with pytest.raises(ValueError, match=sheet + "!"):
        write_workbook_export(path, export, target_environment="local")
    assert not path.exists()


def test_excel_row_limit_includes_header(export, tmp_path, monkeypatch):
    monkeypatch.setattr("dendroflow.configuration.workbook.writer._LAST_EXCEL_ROW", 2)
    path = tmp_path / "control.xlsx"
    with pytest.raises(ValueError, match="too many records"):
        write_workbook_export(path, export, target_environment="local")
    assert not path.exists()


def test_export_never_overwrites_existing_file(export, tmp_path):
    path = tmp_path / "control.xlsx"
    path.write_bytes(b"keep original")
    with pytest.raises(FileExistsError):
        write_workbook_export(path, export, target_environment="local")
    assert path.read_bytes() == b"keep original"


def test_serialization_failure_creates_no_file(export, tmp_path, monkeypatch):
    def fail_save(self, output):
        output.write(b"incomplete")
        raise OSError("serialization failed")

    monkeypatch.setattr(Workbook, "save", fail_save)
    path = tmp_path / "control.xlsx"
    with pytest.raises(OSError, match="serialization failed"):
        write_workbook_export(path, export, target_environment="local")
    assert not path.exists()


def test_export_writer_does_not_connect_to_database(export, save, monkeypatch):
    def fail_connect(*args, **kwargs):
        pytest.fail("writer attempted database access")

    monkeypatch.setattr("dendroflow.database.connect", fail_connect)
    monkeypatch.setattr("dendroflow.configuration.workbook.source.connect", fail_connect)
    save(export)
